from io import BytesIO

import pytest
from docx import Document
from PIL import Image

from itcj2.apps.prorrogas_tec.services.docx_fill import (
    InvalidDocx, MissingValues, fill_docx, find_placeholders, iter_paragraphs, replace_literals,
)


def _docx(build) -> bytes:
    d = Document()
    build(d)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def _text(data: bytes) -> str:
    return "\n".join(p.text for p in iter_paragraphs(Document(BytesIO(data))))


def _png() -> bytes:
    buf = BytesIO()
    Image.new("RGBA", (40, 12), (0, 0, 0, 255)).save(buf, "PNG")
    return buf.getvalue()


def test_encuentra_marcadores_partidos_entre_runs():  # Review Focus 1
    def build(d):
        p = d.add_paragraph()
        p.add_run("Alumno: {{nom")
        p.add_run("bre}} y {{ no_control }}")
    assert find_placeholders(_docx(build)) == {"nombre", "no_control"}


def test_rellena_marcador_partido_con_el_formato_del_primer_run():  # Review Focus 1
    def build(d):
        p = d.add_paragraph()
        p.add_run("Hola ")
        r = p.add_run("{{nom")
        r.bold = True
        p.add_run("bre}}!")
    out = fill_docx(_docx(build), {"nombre": "ANA"})
    para = Document(BytesIO(out)).paragraphs[0]
    assert para.text == "Hola ANA!"
    assert [r.text for r in para.runs if r.bold] == ["ANA"]


def test_rellena_en_tablas_encabezados_y_pies():  # Review Focus 2
    def build(d):
        d.add_table(rows=1, cols=1).cell(0, 0).text = "{{carrera}}"
        d.sections[0].header.paragraphs[0].text = "{{fecha}}"
        d.sections[0].footer.paragraphs[0].text = "{{folio}}"
    out = fill_docx(_docx(build), {"carrera": "ISC", "fecha": "13/enero/2026", "folio": "PR-00001"})
    text = _text(out)
    assert "{{" not in text and {"ISC", "13/enero/2026", "PR-00001"} <= set(text.split("\n"))


def test_faltan_valores_lanza_missing_values_con_los_nombres():
    data = _docx(lambda d: d.add_paragraph("{{a}} {{b}}"))
    with pytest.raises(MissingValues) as e:
        fill_docx(data, {"a": "1"})
    assert e.value.names == frozenset({"b"})


def test_marcador_de_imagen_inserta_la_imagen_y_quita_el_texto():
    out = fill_docx(_docx(lambda d: d.add_paragraph("{{firma_estudiante}}")), {}, images={"firma_estudiante": _png()})
    doc = Document(BytesIO(out))
    assert doc.paragraphs[0].text == "" and len(doc.inline_shapes) == 1


def test_marcador_de_imagen_con_texto_vacio_lo_borra():
    out = fill_docx(_docx(lambda d: d.add_paragraph("Firma: {{firma_estudiante}}")), {"firma_estudiante": ""})
    assert _text(out) == "Firma: "


def test_bytes_que_no_son_docx_lanzan_invalid_docx():  # Review Focus 3
    with pytest.raises(InvalidDocx):
        find_placeholders(b"%PDF-1.7 no soy un docx")


def test_replace_literals_cuenta_y_cruza_runs():
    def build(d):
        p = d.add_paragraph()
        p.add_run("PAYAN FLO")
        p.add_run("RES ITZEL y PAYAN FLORES ITZEL")
    out, counts = replace_literals(_docx(build), {"PAYAN FLORES ITZEL": "{{nombre}}", "NO EXISTE": "x"})
    assert counts == {"PAYAN FLORES ITZEL": 2, "NO EXISTE": 0}
    assert _text(out) == "{{nombre}} y {{nombre}}"
