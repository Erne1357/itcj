from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from itcj2.apps.prorrogas_tec.services.errors import ServiceError
from itcj2.apps.prorrogas_tec.services.students_service import StudentsService

AHORA = datetime(2026, 11, 2, 10, 0)


def test_recordar_una_vez_al_dia(db_session, aceptada, secretaria, con_plan, avisos):
    con_plan(aceptada, [("1000.00", date(2020, 1, 1))])
    StudentsService.remind(db_session, aceptada.id, secretaria.id, now=AHORA)
    assert [a[0] for a in avisos.calls] == ["recordatorio"]
    with pytest.raises(ServiceError) as e:
        StudentsService.remind(db_session, aceptada.id, secretaria.id, now=AHORA + timedelta(hours=23))
    assert e.value.code == "ya_recordado"
    StudentsService.remind(db_session, aceptada.id, secretaria.id, now=AHORA + timedelta(hours=25))


def test_no_se_recuerda_sin_plan_ni_saldo(db_session, aceptada, make_aceptada, secretaria, con_plan, avisos):
    pagada = con_plan(make_aceptada(), [("10.00", date(2020, 1, 1))], pagos=["10.00"])
    db_session.commit()  # el rollback del servicio al fallar descartaría lo que solo está en flush
    with pytest.raises(ServiceError) as e:
        StudentsService.remind(db_session, aceptada.id, secretaria.id, now=AHORA)
    assert e.value.code == "sin_plan"
    with pytest.raises(ServiceError) as e:
        StudentsService.remind(db_session, pagada.id, secretaria.id, now=AHORA)
    assert e.value.code == "sin_saldo"


def test_excel_de_adeudos_se_explica_solo(db_session, aceptada, make_aceptada, secretaria, con_plan, periodo):
    con_plan(aceptada, [("1000.00", date(2020, 1, 1)), ("500.00", date(2091, 5, 1))], pagos=["200.00"])
    al_dia = con_plan(make_aceptada(), [("500.00", date(2091, 5, 1))])
    data, nombre = StudentsService.export_adeudos(db_session, periodo.id, secretaria.id, now=AHORA)
    ws = load_workbook(BytesIO(data)).active
    assert ws["A1"].value == "Lista de adeudos de prórrogas" and "02/11/2026" in ws["A3"].value
    assert "parcialidad vencida y no cubierta" in ws["A4"].value and ws.freeze_panes == "A8" and ws.auto_filter.ref
    assert [c.value for c in ws[7]][:3] == ["No. de control", "Nombre", "Carrera"]
    controles = [r[0].value for r in ws.iter_rows(min_row=8)]
    assert aceptada.control_number in controles and al_dia.control_number not in controles
    assert nombre.startswith("adeudos-prorrogas-") and nombre.endswith(".xlsx")
    primera = aceptada.block_exported_at
    StudentsService.export_adeudos(db_session, periodo.id, secretaria.id, now=AHORA + timedelta(days=3))
    db_session.refresh(aceptada)
    assert primera is not None and aceptada.block_exported_at == primera


def test_avisar_desbloqueo_solo_si_aplica(db_session, aceptada, secretaria, con_plan, periodo):
    con_plan(aceptada, [("1000.00", date(2020, 1, 1))])
    StudentsService.export_adeudos(db_session, periodo.id, secretaria.id, now=AHORA)
    with pytest.raises(ServiceError) as e:
        StudentsService.mark_unblock(db_session, aceptada.id, secretaria.id, now=AHORA)
    assert e.value.code == "no_aplica"                       # exportada pero aún debe
    from itcj2.apps.prorrogas_tec.models.plan import ProrrogaPayment
    db_session.add(ProrrogaPayment(request_id=aceptada.id, folio="Z-1", paid_on=date(2026, 11, 1),
                                   amount=1000))
    db_session.flush()
    StudentsService.mark_unblock(db_session, aceptada.id, secretaria.id, now=AHORA)
    assert aceptada.unblock_notified_at is not None


def test_api_exportar(client_as, secretaria, estudiante, periodo):
    r = client_as(secretaria).post("/api/prorrogas/v2/estudiantes/adeudos.xlsx", json={"period_id": periodo.id})
    assert r.status_code == 200 and r.content[:2] == b"PK" and "attachment" in r.headers["content-disposition"]
    assert client_as(estudiante).post("/api/prorrogas/v2/estudiantes/adeudos.xlsx",
                                      json={"period_id": periodo.id}).status_code == 403


def test_reexportar_tras_avisar_desbloqueo_reinicia_el_ciclo(db_session, aceptada, secretaria, con_plan, periodo):  # I1
    from itcj2.apps.prorrogas_tec.models.event import ProrrogaEvent
    from itcj2.apps.prorrogas_tec.services.payment_service import PaymentService

    def pagar(folio):
        return PaymentService.record(db_session, aceptada.id, secretaria.id, folio=folio,
                                     paid_on=aceptada.created_at.date(), amount=Decimal("1000"),
                                     now=aceptada.created_at)

    con_plan(aceptada, [("1000.00", date(2020, 1, 1))])
    StudentsService.export_adeudos(db_session, periodo.id, secretaria.id, now=AHORA)
    pago = pagar(f"I1-{aceptada.id}-a")
    StudentsService.mark_unblock(db_session, aceptada.id, secretaria.id, now=AHORA)
    PaymentService.void(db_session, pago.id, secretaria.id, motivo="Recibo de otro alumno", now=AHORA)
    segunda = AHORA + timedelta(days=3)
    StudentsService.export_adeudos(db_session, periodo.id, secretaria.id, now=segunda)
    db_session.refresh(aceptada)
    assert aceptada.unblock_notified_at is None and aceptada.block_exported_at == segunda
    ev = (db_session.query(ProrrogaEvent).filter_by(event_type="lista_bloqueo_exportada")
          .order_by(ProrrogaEvent.id.desc()).first())
    assert ev.payload["solicitudes"] == [aceptada.id] and ev.payload["estudiantes"] == 1
    pagar(f"I1-{aceptada.id}-b")
    assert StudentsService.expediente(db_session, aceptada.id, segunda.date())["bloqueo"]["avisar_desbloqueo"] is True


def test_excel_no_convierte_texto_en_formula(db_session, aceptada, secretaria, con_plan, periodo):  # T7
    aceptada.student_name = "=1+1"
    con_plan(aceptada, [("1000.00", date(2020, 1, 1))])
    data, _ = StudentsService.export_adeudos(db_session, periodo.id, secretaria.id, now=AHORA)
    ws = load_workbook(BytesIO(data)).active
    fila = next(r for r in ws.iter_rows(min_row=8) if r[0].value == aceptada.control_number)
    assert fila[1].value == "=1+1" and fila[1].data_type == "s"


def test_recordatorio_dice_si_el_pago_ya_vencio(db_session, aceptada, make_aceptada, secretaria, con_plan,
                                                enviar_correo):  # M9
    from itcj2.core.models.notification import Notification
    from itcj2.templates import templates

    con_plan(aceptada, [("1000.00", date(2020, 1, 1))])
    al_dia = con_plan(make_aceptada(), [("500.00", date(2091, 5, 1))])
    for req in (aceptada, al_dia):
        StudentsService.remind(db_session, req.id, secretaria.id, now=AHORA)
    cuerpos = [db_session.query(Notification).filter_by(user_id=r.student_id, type="prorroga_recordatorio").one().body
               for r in (aceptada, al_dia)]
    assert cuerpos == ["Saldo $1,000.00. Tienes un pago vencido desde el 01 de enero de 2020 por $1,000.00.",
                       "Saldo $500.00. Próximo pago: 01 de mayo de 2091 por $500.00."]
    plantilla = templates.env.get_template("prorrogas_tec/email/recordatorio.html")
    vencido, proximo = (plantilla.render(**c["context"]) for c in enviar_correo.calls)
    assert "pago vencido desde el 01 de enero de 2020 por $1,000.00" in vencido and "próximo pago" not in vencido
    assert "Tu próximo pago es el 01 de mayo de 2091 por $500.00" in proximo and "vencido" not in proximo
