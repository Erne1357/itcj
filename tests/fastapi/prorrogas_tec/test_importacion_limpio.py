from datetime import date
from decimal import Decimal as D

import pytest
from openpyxl import Workbook, load_workbook

from itcj2.apps.prorrogas_tec.importacion.hojas import Parcialidad, Solicitud, leer_libro
from itcj2.apps.prorrogas_tec.importacion.limpio import ENCABEZADOS, escribir_limpio, leer_limpio

ENC = ["No.", "FECHA", "No. De Control", "NOMBRE", "CARRERA", "ENERO", "IMPORTE", "", "FEBRERO", "IMPORTE", "",
       "MARZO", "IMPORTE", ""]


def _origen(libro):
    return libro({"Inscripción Enero - Junio 2026": [ENC,
        [None, "12 de enero de 2026", 99120001, "Buena", "Contador Público", "16 de enero de 2026", 1000, 1,
         "06 de febrero de 2026", 1000, 0],
        [None, "12 de enero de 2026", 99120002, "Mala", "Contador Público", "16 de enero de 2026", 1000, 1,
         "06 de febrero de 2026", "Prórroga cancelada", 1, "06 de marzo de 2026", 1000, 1]]})


def _sol(parcs, ref="H#2", fecha=date(2026, 1, 12), problema="", control="99120001", nombre="Ana"):
    return Solicitud(ref, "H", "enero – junio 2026", fecha, control, nombre, "Contador Público", "CP",
                     parcs, problema)


def _p(dia, monto=D("1000"), pagado=True):
    return Parcialidad(date(2026, 2, dia), monto, pagado, origen="H!F2")


def test_ida_y_vuelta(libro, tmp_path):
    salida = tmp_path / "limpio.xlsx"
    res = escribir_limpio(leer_libro(_origen(libro)), salida)
    assert res == {"solicitudes": 2, "parcialidades": 5, "con_problema": 1}
    wb = load_workbook(salida)
    assert [c.value for c in wb["Parcialidades"][1]] == list(ENCABEZADOS) and wb["Revisar"].max_row == 2
    validas, errores = leer_limpio(salida)
    buena, = validas
    assert buena.no_control == "99120001" and buena.parcialidades == [
        (1, date(2026, 1, 16), D("1000.00"), True), (2, date(2026, 2, 6), D("1000.00"), False)]
    assert [r for r, _ in errores] == ["Inscripción Enero - Junio 2026#3"]


def test_problema_en_una_parcialidad_marca_toda_la_solicitud(libro, tmp_path):  # Review Focus 2
    salida = tmp_path / "limpio.xlsx"
    escribir_limpio(leer_libro(_origen(libro)), salida)
    filas = [r for r in load_workbook(salida)["Parcialidades"].iter_rows(min_row=2, values_only=True)
             if r[0].endswith("#3")]
    assert len(filas) == 3 and sum(1 for r in filas if r[ENCABEZADOS.index("problema")]) == 1
    _, errores = leer_limpio(salida)
    assert errores[0][1].startswith("tiene datos por revisar")


def test_encabezados_cambiados_se_rechazan(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "Parcialidades"; ws.append(["ref", "otra"]); wb.save(tmp_path / "x.xlsx")
    with pytest.raises(ValueError):
        leer_limpio(tmp_path / "x.xlsx")


def test_sin_parcialidades(tmp_path):
    salida = tmp_path / "l.xlsx"
    res = escribir_limpio([_sol([])], salida)
    assert res == {"solicitudes": 1, "parcialidades": 0, "con_problema": 1}
    assert leer_limpio(salida) == ([], [("H#2", "tiene datos por revisar: sin parcialidades")])


def test_sin_fecha_de_solicitud(tmp_path):
    salida = tmp_path / "l.xlsx"
    escribir_limpio([_sol([_p(5)], fecha=None)], salida)
    assert leer_limpio(salida) == ([], [("H#2", "sin fecha de solicitud")])


def test_numeracion_por_vencimiento(tmp_path):
    salida = tmp_path / "l.xlsx"
    escribir_limpio([_sol([_p(20, D("3")), _p(5, D("1")), _p(5, D("2"))])], salida)
    validas, errores = leer_limpio(salida)
    assert not errores
    assert [(n, f.day, m) for n, f, m, _ in validas[0].parcialidades] == [
        (1, 5, D("1.00")), (2, 5, D("2.00")), (3, 20, D("3.00"))]


def test_alumno_distinto_entre_renglones(tmp_path):
    salida = tmp_path / "l.xlsx"
    escribir_limpio([_sol([_p(5), _p(6)])], salida)
    wb = load_workbook(salida)
    wb["Parcialidades"].cell(3, ENCABEZADOS.index("nombre") + 1).value = "Otra"
    wb.save(salida)
    assert leer_limpio(salida) == ([], [("H#2", "datos del estudiante distintos entre renglones")])


def test_problema_de_solicitud_en_todas_las_filas_y_revisar(tmp_path):
    salida = tmp_path / "l.xlsx"
    p = Parcialidad(None, None, None, "fecha ilegible", "xx", origen="H!F2")
    escribir_limpio([_sol([_p(5), p], problema="carrera sin equivalencia")], salida)
    wb = load_workbook(salida)
    filas = list(wb["Parcialidades"].iter_rows(min_row=2, values_only=True))
    pr = ENCABEZADOS.index("problema")
    assert [r[pr] for r in filas] == ["carrera sin equivalencia", "carrera sin equivalencia; fecha ilegible"]
    assert [r[ENCABEZADOS.index("numero")] for r in filas] == [1, 2]
    rev = list(wb["Revisar"].iter_rows(min_row=2, values_only=True))
    assert rev[0][:2] == ("H", 2) or rev[0][:2] == ("H", "2")
    assert len(rev) == 2


def test_cli_excel_limpiar_no_imprime_datos(libro, tmp_path):
    from click.testing import CliRunner
    from itcj2.cli.prorrogas import prorrogas_cli
    salida = tmp_path / "out.xlsx"
    r = CliRunner().invoke(prorrogas_cli, ["excel-limpiar", str(_origen(libro)), "--salida", str(salida)])
    assert r.exit_code == 0, r.output
    assert "Solicitudes: 2 · Parcialidades: 5 · Con problema: 1" in r.output and str(salida) in r.output
    assert "Mala" not in r.output and "99120002" not in r.output
