from io import BytesIO

import pytest
from docx import Document

from itcj2.apps.prorrogas_tec.documents.catalog import DOC_TYPES, default_template_path, sample_values
from itcj2.apps.prorrogas_tec.services.docx_fill import fill_docx, find_placeholders, iter_paragraphs

PII = ["PAYAN", "ITZEL", "22111670", "HERNANDEZ SOLANO", "24112010", "CARDENAS ALMANZA", "M23111966"]


@pytest.mark.parametrize("code", list(DOC_TYPES))
def test_plantilla_del_sistema_tiene_los_marcadores_correctos(code):
    found = find_placeholders(default_template_path(code).read_bytes())
    assert DOC_TYPES[code].required <= found <= DOC_TYPES[code].allowed


@pytest.mark.parametrize("code", list(DOC_TYPES))
def test_plantilla_del_sistema_no_trae_datos_de_estudiantes_reales(code):
    doc = Document(BytesIO(default_template_path(code).read_bytes()))
    text = "\n".join(p.text for p in iter_paragraphs(doc)).upper()
    assert not [s for s in PII if s in text]


@pytest.mark.parametrize("code", list(DOC_TYPES))
def test_plantilla_del_sistema_se_llena_con_los_datos_de_ejemplo(code):
    out = fill_docx(default_template_path(code).read_bytes(), *sample_values(code))
    assert find_placeholders(out) == set()
