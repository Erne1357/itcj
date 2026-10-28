"""Servicio del reporte Excel de la encuesta de egresados (solo lectura).

Formulario sintetico con codigo unico por prueba: la BD de dev trae el real
`egresados`, y el indice parcial deja una sola version abierta por `code`.
"""
import io
import uuid
from datetime import datetime

import pytest

from itcj2.core.utils.timezone import db_now

SIN_RESP = "Sin respuesta"

SCHEMA = {
    "sections": [{"key": "perfil", "title": "Perfil"}, {"key": "otras", "title": "Otras"}],
    "fields": [
        {"key": "nombre_completo", "type": "text", "label": "Nombre completo (con apellidos):",
         "section": "perfil"},
        {"key": "carrera_egreso", "type": "radio", "label": "Carrera de egreso: (escoja una)",
         "section": "perfil",
         "options": [{"label": "Sistemas", "value": "Sistemas"},
                     {"label": "Electrónica", "value": "Electrónica"},
                     {"label": "Industrial", "value": "Industrial"}]},
        {"key": "recibir_correos", "type": "radio", "label": "¿Recibir correos?",
         "section": "perfil",
         "options": [{"label": "Si", "value": "Si"}, {"label": "No", "value": "No"}]},
        {"key": "anio_egreso", "type": "text", "label": "Año de egreso", "section": "perfil"},
        {"key": "idiomas", "type": "multiselect", "label": "Idiomas", "section": "otras",
         "options": [{"label": "Inglés", "value": "en"}, {"label": "Francés", "value": "fr"}]},
        {"key": "scale_x", "type": "scale", "label": "Titulado", "section": "otras",
         "scale": {"min": 1, "max": 5}},
        {"key": "fecha_nacimiento", "type": "date", "label": "Nacimiento", "section": "otras"},
        {"key": "comentario", "type": "textarea", "label": "Comentario", "section": "otras"},
    ],
}


@pytest.fixture()
def form(db_session):
    from itcj2.apps.titulatec.models import SurveyForm
    f = SurveyForm(code=f"rep_{uuid.uuid4().hex[:8]}", title="Encuesta de prueba",
                   schema=SCHEMA, version=1, status="closed", is_anonymous=False)
    db_session.add(f)
    db_session.flush()
    return f


@pytest.fixture()
def resp(db_session, form):
    from itcj2.apps.titulatec.models import SurveyResponse

    def _make(answers, *, fecha=(2026, 3, 1), control=None, source="session", user=None):
        r = SurveyResponse(form_id=form.id, form_version=1, identity_source=source,
                           answers=answers, control_number=control,
                           user_id=user.id if user else None,
                           submitted_at=datetime(*fecha, 10, 0))
        db_session.add(r)
        db_session.flush()
        return r
    return _make


def S():
    from itcj2.apps.titulatec.services.survey_report_service import SurveyReportService
    return SurveyReportService


def P(form, **q):
    return S().parse_params(q, form)


def collect(db, form, **q):
    return S().collect(db, form, P(form, **q))


def test_parse_y_to_query_ida_y_vuelta(form):
    p = P(form, cols="_nombre,carrera_egreso,basura", orden="carrera_egreso:desc,_nombre",
          desde="2026-01-01", hasta="2026-06-30", egreso_desde="2010",
          f_carrera_egreso=["Sistemas", "Industrial"], f_recibir_correos="Si",
          unico="1", encabezado="clave", hojas="respuestas,resumen")
    assert p.cols == ("_nombre", "carrera_egreso")
    assert any("basura" in w for w in p.warnings)
    assert S().parse_params(S().to_query(p), form) == p
    from starlette.datastructures import QueryParams
    assert S().parse_params(QueryParams(S().to_query(p)), form) == p
    d = P(form)
    assert d.hojas == ("respuestas", "resumen", "distribucion", "cruce", "parametros")
    assert S().parse_params(S().to_query(d), form) == d


def test_fecha_invalida_se_ignora_con_aviso_y_hojas_presente_es_exacta(form):
    p = P(form, desde="ayer", hojas="respuestas")
    assert p.desde is None and any("desde" in w for w in p.warnings)
    assert p.hojas == ("respuestas",)
    assert P(form, hojas="resumen").hojas == ("respuestas", "resumen")


def test_filtros_fecha_origen_y_opcion(db_session, form, resp):
    resp({"carrera_egreso": "Sistemas", "recibir_correos": "Si"}, fecha=(2026, 1, 10))
    resp({"carrera_egreso": "Industrial", "recibir_correos": "No"}, fecha=(2026, 6, 30))
    resp({"carrera_egreso": "Sistemas", "recibir_correos": "No"}, fecha=(2026, 8, 1),
         source="import")
    assert len(collect(db_session, form)) == 3
    # `hasta` es inclusivo hasta el fin del dia (la respuesta es 30/06 10:00).
    assert len(collect(db_session, form, hasta="2026-06-30")) == 2
    assert len(collect(db_session, form, desde="2026-07-01")) == 1
    assert [f["_origen"] for f in collect(db_session, form, origen="importada")] == ["Importada"]
    assert len(collect(db_session, form, origen="plataforma")) == 2
    # Valores de un mismo f_ = OR; varios f_ = AND.
    assert len(collect(db_session, form, f_carrera_egreso=["Sistemas", "Industrial"])) == 3
    assert len(collect(db_session, form, f_carrera_egreso="Sistemas",
                       f_recibir_correos="No")) == 1


def test_vacio_y_filtro_inexistente(db_session, form, resp):
    resp({"recibir_correos": "Si"})
    resp({})
    assert len(collect(db_session, form, f_recibir_correos="__vacio__")) == 1
    assert len(collect(db_session, form, f_recibir_correos=["Si", "__vacio__"])) == 2
    p = P(form, f_nombre_completo="x")          # no es de opcion: se descarta
    assert p.filtros == () and p.warnings


def test_multiselect_y_scale_en_filtros(db_session, form, resp):
    resp({"idiomas": ["en", "fr"], "scale_x": 5})
    resp({"idiomas": ["fr"], "scale_x": 3})
    resp({"idiomas": [], "scale_x": None})
    assert len(collect(db_session, form, f_idiomas="en")) == 1       # alguno de la lista
    assert len(collect(db_session, form, f_idiomas="fr")) == 2
    assert len(collect(db_session, form, f_scale_x="5")) == 1
    assert len(collect(db_session, form, f_idiomas="__vacio__")) == 1


def test_rango_de_egreso_deja_fuera_al_que_no_tiene_anio_valido(db_session, form, resp):
    resp({"anio_egreso": "2012"})
    resp({"anio_egreso": "2020"})
    resp({"anio_egreso": "dos mil"})
    resp({})
    assert len(collect(db_session, form, egreso_desde="2010", egreso_hasta="2015")) == 1
    assert len(collect(db_session, form, egreso_desde="2010")) == 2
    assert len(collect(db_session, form)) == 4


def test_unico_deja_la_mas_reciente_y_conserva_las_sin_control(db_session, form, resp):
    resp({"recibir_correos": "Si"}, fecha=(2026, 1, 1), control="99000001")
    resp({"recibir_correos": "No"}, fecha=(2026, 2, 1), control="99000001")
    resp({"recibir_correos": "Si"}, fecha=(2026, 1, 5))
    resp({"recibir_correos": "Si"}, fecha=(2026, 1, 6))
    assert len(collect(db_session, form, unico="1")) == 3
    # Quien dijo «Si» y luego «No» no se cuela en el filtro «Si».
    assert len(collect(db_session, form, unico="1", f_recibir_correos="Si")) == 2
    assert S()._unicos(collect(db_session, form)) == 3


def test_orden_vacios_al_final_acentos_y_numeros(db_session, form, resp):
    resp({"nombre_completo": "Zeta", "anio_egreso": "2009"})
    resp({"nombre_completo": "Álvaro", "anio_egreso": "2010"})
    resp({"nombre_completo": "Beto", "anio_egreso": "999"})
    resp({})
    c = collect(db_session, form, orden="nombre_completo:asc", cols="nombre_completo")
    assert [f["nombre_completo"] for f in c] == ["Álvaro", "Beto", "Zeta", None]
    c = collect(db_session, form, orden="nombre_completo:desc", cols="nombre_completo")
    assert [f["nombre_completo"] for f in c] == ["Zeta", "Beto", "Álvaro", None]  # vacio al final aun en desc
    # Numeros como numero: 999 < 2009 < 2010 (como texto 2009 < 2010 < 999).
    c = collect(db_session, form, orden="anio_egreso:asc", cols="anio_egreso")
    assert [f["anio_egreso"] for f in c] == [999, 2009, 2010, None]


def test_presets_se_adaptan_y_catalogo(form):
    cat = S().catalog(form)
    assert {p["key"] for p in cat["presets"]} == {
        "seguimiento", "perfil", "laboral", "satisfaccion", "completo"}
    seg = next(p for p in cat["presets"] if p["key"] == "seguimiento")
    assert "especialidad" not in seg["cols"] and "carrera_egreso" in seg["cols"]
    assert "f_recibir_correos=Si" in seg["query"] and "unico=1" in seg["query"]
    assert set(cat["choice_keys"]) == {"carrera_egreso", "recibir_correos", "idiomas", "scale_x"}
    assert [c["key"] for c in cat["meta"]][:3] == ["_id", "_fecha", "_origen"]
    assert len(cat["meta"]) == 9
    todas = [c for s in cat["sections"] for c in s["columns"]]
    carrera = next(c for c in todas if c["key"] == "carrera_egreso")
    assert carrera["short"] == "Carrera" and carrera["choice"]
    assert carrera["options"][0] == {"value": "Sistemas", "label": "Sistemas"}
    escala = next(c for c in todas if c["key"] == "scale_x")
    assert escala["short"] == "Valora: Titulado" and len(escala["options"]) == 5
    # La query del preset reparsea al mismo preset; sin cols cae en `seguimiento`.
    assert S().to_query(S().parse_params(seg["query"], form)) == seg["query"]
    assert P(form).cols == tuple(seg["cols"])


def test_semesters(db_session, form, resp):
    resp({}, fecha=(2026, 2, 1))
    resp({}, fecha=(2026, 9, 1))
    resp({}, fecha=(2025, 12, 31))
    s = S().semesters(db_session, form)
    assert [x["code"] for x in s] == ["2026B", "2026A", "2025B"]
    assert s[1]["desde"].isoformat() == "2026-01-01"
    assert s[1]["hasta"].isoformat() == "2026-06-30"


def _libro(db, form, **q):
    from openpyxl import load_workbook
    p = P(form, **q)
    data = S().build_xlsx(db, form, p, generated_by="Quien Sea", generated_at=db_now())
    assert data[:2] == b"PK"
    return load_workbook(io.BytesIO(data)), p


def test_libro_hojas_encabezados_y_valores_tipados(db_session, form, resp, make_user):
    u = make_user(first_name="ANA", last_name="PEREZ", control_number="99000007")
    resp({"carrera_egreso": "Sistemas", "recibir_correos": "Si", "anio_egreso": "2015",
          "idiomas": ["en", "fr"], "scale_x": 4, "fecha_nacimiento": "1990-05-17",
          "comentario": "=CMD()"}, user=u)
    resp({"carrera_egreso": "Industrial", "recibir_correos": "No"}, source="import")
    cols = ("_id,_fecha,_control,_nombre,carrera_egreso,anio_egreso,idiomas,scale_x,"
            "fecha_nacimiento,comentario")
    wb, _ = _libro(db_session, form, cols=cols, encabezado="clave")
    assert wb.sheetnames == ["Respuestas", "Resumen", "Distribución", "Cruce", "Parámetros"]
    ws = wb["Respuestas"]
    assert ws["A1"].value.endswith("Reporte")
    assert "2 respuestas" in ws["A2"].value and "Quien Sea" in ws["A3"].value
    enc = [c.value for c in ws[5]]
    assert enc == cols.split(",")
    fila = {enc[i]: ws.cell(row=6, column=i + 1) for i in range(len(enc))}
    assert isinstance(fila["_fecha"].value, datetime)
    assert isinstance(fila["_id"].value, (int, float))
    assert fila["anio_egreso"].value == 2015
    assert fila["scale_x"].value == 4
    assert fila["_control"].value == "99000007"
    assert fila["_nombre"].value == "PEREZ ANA"
    assert fila["idiomas"].value == "en; fr"
    assert isinstance(fila["fecha_nacimiento"].value, datetime)
    # `=CMD()` queda como texto, nunca como formula.
    assert fila["comentario"].value == "=CMD()" and fila["comentario"].data_type == "s"
    assert ws.freeze_panes == "A6"          # _control/_nombre no son de las 2 primeras


def test_libro_congela_control_y_nombre_si_van_al_inicio(db_session, form, resp):
    resp({"carrera_egreso": "Sistemas"})
    wb, _ = _libro(db_session, form, cols="_control,_nombre,carrera_egreso", hojas="respuestas")
    assert wb["Respuestas"].freeze_panes == "C6"


def test_libro_texto_con_formula_en_cualquier_hoja_es_texto(db_session, form, resp):
    resp({"nombre_completo": "=1+1", "carrera_egreso": "+SUM(A1)"})
    wb, _ = _libro(db_session, form, cols="nombre_completo,carrera_egreso",
                   encabezado="clave")
    ws = wb["Respuestas"]
    assert [c.data_type for c in ws[6]] == ["s", "s"]
    assert ws["A6"].value == "=1+1"
    # El valor fuera de catalogo («+SUM(A1)») aparece en Distribucion y Resumen como texto.
    for hoja in ("Distribución", "Resumen", "Cruce"):
        assert all(c.data_type != "f" for row in wb[hoja].iter_rows() for c in row)


def test_libro_solo_respuestas_y_sin_filas(db_session, form):
    wb, _ = _libro(db_session, form, hojas="respuestas", cols="carrera_egreso")
    assert wb.sheetnames == ["Respuestas"]
    assert wb["Respuestas"]["A6"].value == "Sin respuestas con estos filtros"
    wb, _ = _libro(db_session, form)            # todas las hojas, cero filas: no revienta
    assert len(wb.sheetnames) == 5


def test_libro_cruce_con_llave_invalida_da_aviso_no_error(db_session, form, resp):
    resp({"carrera_egreso": "Sistemas", "recibir_correos": "Si"})
    wb, p = _libro(db_session, form, cruce_filas="nombre_completo")
    # parse la reemplaza por el default y avisa...
    assert any("nombre_completo" in w for w in p.warnings)
    # ...y si llega una llave invalida por otro camino, la hoja trae aviso (no error).
    from dataclasses import replace

    from openpyxl import load_workbook
    forzado = replace(p, cruce_filas="nombre_completo")
    data = S().build_xlsx(db_session, form, forzado, generated_by="x", generated_at=db_now())
    ws = load_workbook(io.BytesIO(data))["Cruce"]
    assert "No se pudo armar el cruce" in ws["A3"].value


def test_libro_distribucion_multiselect_scale_y_cruce(db_session, form, resp):
    resp({"carrera_egreso": "Sistemas", "recibir_correos": "Si",
          "idiomas": ["en", "fr"], "scale_x": 5})
    resp({"carrera_egreso": "Sistemas", "recibir_correos": "No",
          "idiomas": ["fr"], "scale_x": 3})
    resp({"carrera_egreso": "Rara", "scale_x": 3}, source="import")
    wb, _ = _libro(db_session, form, cols="carrera_egreso,idiomas,scale_x")
    texto = [c.value for row in wb["Distribución"].iter_rows() for c in row
             if c.value is not None]
    assert "Promedio" in texto and "Rara" in texto and SIN_RESP in texto
    assert any("más de una opción" in str(t) for t in texto)
    cruce = [c.value for row in wb["Cruce"].iter_rows() for c in row if c.value is not None]
    assert "Sistemas" in cruce and "Total" in cruce


def test_preview_y_filename(db_session, form, resp):
    for i in range(12):
        resp({"carrera_egreso": "Sistemas", "recibir_correos": "Si"}, fecha=(2026, 1, 1 + i))
    pv = S().preview(db_session, form,
                     P(form, cols="_fecha,carrera_egreso", encabezado="pregunta"), limit=10)
    assert pv["total"] == 12 and pv["unique"] == 12 and len(pv["rows"]) == 10
    assert pv["columns"][1] == {"key": "carrera_egreso",
                                "header": "Carrera de egreso: (escoja una)",
                                "full": "Carrera de egreso: (escoja una)"}
    assert pv["rows"][0] == ["01/01/2026 10:00", "Sistemas"]
    assert "cols=_fecha%2Ccarrera_egreso" in pv["query"] and pv["warnings"] == []
    assert S().filename(form, datetime(2026, 10, 28, 9, 5)) == \
        f"reporte-encuesta-{form.code}-v1-20261028-0905.xlsx"


def test_headers_repetidos_se_vuelven_unicos():
    from itcj2.apps.titulatec.services.survey_report_service import _unicos_encabezados
    assert _unicos_encabezados(["A", "a", "B", "A"]) == ["A", "a (2)", "B", "A (3)"]
