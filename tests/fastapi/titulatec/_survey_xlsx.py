"""Fixture SINTÉTICO del Excel de Microsoft Forms de la encuesta de egresados
(spec `2026-10-05-titulatec-import-encuesta-xlsx-design.md` §2/§4.2).

`HEADERS` es una copia LITERAL de la fila 1 del archivo real (69 columnas,
solo el texto de los encabezados, que no tiene datos personales: saltos de
línea y los espacios duros (`\\xa0`) de «INGRESO», «contratación» y «empresa:» son
del original y se conservan a propósito). Todo lo demás -nombres, controles,
respuestas- es inventado: el archivo real NUNCA se usa en pruebas.

`SCHEMA` es un esquema sintético con los mismos `key`/`type` del formulario
`egresados` v1 (opciones recortadas a lo que ejercen las pruebas).
"""
from __future__ import annotations

import io
from datetime import datetime

ASPECTO = ("Aspecto que valora la empresa u organismo para la contratación de "
           "egresados.")

HEADERS = [
    "Id",
    "Start time",
    "Completion time",
    "Email",
    "Name",
    "Nombre (s) y Apellidos completos: (Sino escribe de forma correcta su información "
    "personal y/o adicional, tendrá que llenar de nuevo la encuesta... Sin excepción de "
    "Casos).\n",
    "No. Control:",
    "Fecha de nacimiento:",
    "Sexo",
    "Estado civil",
    "Correo Personal: (Cuenta electrónica que sea utilizada de forma continua, no se "
    "reenviaran correos)",
    "Número telefónico:",
    "Carrera de Egreso: De no encontrar su carrera, escoger la que más se apegue a su "
    "profesión cursada.",
    "Especialidad:",
    "Periodo de INGRESO",
    "Año de INGRESO\n\xa0(EJ. 1999)",
    "Periodo de EGRESO",
    "Año de EGRESO (EJ. 1999)",
    "Promedio final obtenido (Ej.93)",
    "En que año realizo las Residencias (Ej. 2000)",
    "Le interesaría recibir correos con información relevante para usted (Ej. Ofertas "
    "de empleo, cursos ofertados en el ITCJ, etc..)",
    "Calidad de los docentes:",
    "Plan de Estudios:",
    "Oportunidad de participar en proyectos de investigación y desarrollo:",
    "Énfasis que se le prestaba a la investigación dentro del proceso de enseñanza:",
    "Satisfacción con las condiciones de estudio (infraestructura)",
    "Experiencia obtenida a través de la Residencia Profesional",
    "¿Cómo acreditó el requisito del segundo idioma para titulación? (inglés):",
    "¿Domina algún idioma distinto a su lengua materna?:",
    "¿Qué idioma?",
    "Actividad a la que se dedica actualmente:",
    "Si se encuentra estudiando, seleccione la opción que más se apegue a la actividad "
    "que realiza (estudios):",
    "Nombre de la institución (Escuela) dónde estudia: en caso de no estudiar, poner No "
    "Estudio.",
    "En caso de trabajar, tiempo transcurrido para obtener el primer empleo:",
    "Medios para obtener el empleo",
    "Requisito de contratación\xa0",
    "¿Idioma qué más utiliza en su trabajo? Aparte de su lengua materna:",
    "Que habilidad del idioma extranjero utiliza con mayor frecuencia en sus actividades "
    "laborales",
    "Nivel Jerárquico en el trabajo:",
    "Condición de trabajo (compromiso):",
    "Las actividades que realiza en su trabajo ¿Están relacionado con su carrera de "
    "egreso?:",
    "Funciones que realiza en la empresa:\xa0 (si no trabaja, poner \"no trabajo\")",
    "Si trabaja, escriba el nombre de la empresa (sino trabaja... solo responder \"No "
    "trabajo\")",
    "La empresa u organismo donde labora es:",
    "La empresa donde labora pertenece al sector:",
    "Nombre Completo del encargado de Recursos Humanos:",
    "Teléfono de contacto del encargado de Recursos Humanos:",
    "Tamaño de la empresa u organización",
    "Antigüedad en el empleo actual",
    "Cuenta con alguna empresa propia:",
    "Nombre de la empresa: (en caso de no tener, escriba ¨N0¨)",
    "Cómo califica su formación académica con respecto a su desempeño laboral:",
    "Utilidad de las residencias profesionales o prácticas profesionales para su "
    "desarrollo laboral y profesional",
    ASPECTO + "Área o campo de estudio",
    ASPECTO + "Titulado",
    ASPECTO + "Experiencia laboral/practica (antes de egresar)",
    ASPECTO + "Competencia laboral: Ej. Habilidad para resolver problemas, etc..",
    ASPECTO + "Institución de egreso",
    ASPECTO + "Dominio de otro idioma",
    ASPECTO + "Recomendación",
    ASPECTO + "Personalidad/Actitud",
    ASPECTO + "Capacidad de liderazgo",
    ASPECTO + "No trabajo",
    "Pertenece a organizaciones sociales",
    "Nombre de la organización social",
    "Pertenece a alguna asociación de egresados:",
    "Nombre de la asociación de egresados",
    "Le interesaría pertenecer a la asociación de egresados del ITCJ",
    "Comentario, queja y/o sugerencia que desee aportar, para mejorar el sistema de la "
    "organización... Es de gran importancia su respuesta (evite las respuestas NO, NADA, "
    "PUNTO... no son válidas).",
]

# `field_key` esperado por columna (None = columnas de metadatos A..E).
KEYS = [None, None, None, None, None, "nombre_completo", "no_control",
        "fecha_nacimiento", "sexo", "estado_civil", "correo_personal", "telefono",
        "carrera_egreso", "especialidad", "periodo_ingreso", "anio_ingreso",
        "periodo_egreso", "anio_egreso", "promedio_final", "anio_residencias",
        "recibir_correos", "calidad_docentes", "plan_estudios",
        "oportunidad_investigacion", "enfasis_investigacion",
        "satisfaccion_infraestructura", "experiencia_residencia", "acreditacion_idioma",
        "domina_otro_idioma", "que_idioma", "actividad_actual", "tipo_estudio",
        "institucion_estudio", "tiempo_primer_empleo", "medio_obtencion_empleo",
        "requisito_contratacion", "idioma_uso_trabajo", "habilidad_idioma_trabajo",
        "nivel_jerarquico", "condicion_trabajo", "relacion_actividad_carrera",
        "funciones_empresa", "nombre_empresa", "tipo_organismo", "sector_empresa",
        "nombre_encargado_rh", "telefono_encargado_rh", "tamano_empresa",
        "antiguedad_empleo", "empresa_propia", "nombre_empresa_propia",
        "calif_formacion_academica", "utilidad_residencias", "scale_area_estudio",
        "scale_titulado", "scale_experiencia_laboral", "scale_competencia_laboral",
        "scale_institucion_egreso", "scale_dominio_idioma", "scale_recomendacion",
        "scale_personalidad_actitud", "scale_capacidad_liderazgo",
        "extra_aspecto_no_trabajo", "pertenece_org_social", "nombre_org_social",
        "pertenece_asoc_egresados", "nombre_asoc_egresados", "interes_asoc_egresados",
        "comentario_sugerencia"]
assert len(HEADERS) == len(KEYS) == 69

_TEXT = {"nombre_completo": 200, "correo_personal": 150, "telefono": 20,
         "institucion_estudio": 200, "funciones_empresa": 500, "nombre_empresa": 200,
         "nombre_encargado_rh": 160, "telefono_encargado_rh": 20,
         "nombre_empresa_propia": 200, "nombre_org_social": 200,
         "nombre_asoc_egresados": 200}
_OPCIONES = {
    "sexo": ["Hombre", "Mujer"],
    "estado_civil": ["Soltero (a)", "Casado (a)", "Unión libre"],
    "actividad_actual": ["Estudia", "Trabaja", "Estudia y trabaja",
                         "No estudia, ni trabaja"],
    "periodo_ingreso": ["AGOSTO DICIEMBRE", "ENERO JUNIO"],
    "periodo_egreso": ["AGOSTO DICIEMBRE", "ENERO JUNIO"],
    "acreditacion_idioma": ["Examen de Titulación", "Aprobó nivel III"],
}


def _field(key: str) -> dict:
    if key == "no_control":
        return {"key": key, "type": "text",
                "validation": {"format": "digits", "length": 8, "maxLength": 8}}
    if key == "fecha_nacimiento":
        return {"key": key, "type": "date"}
    if key in ("anio_ingreso", "anio_egreso", "anio_residencias"):
        return {"key": key, "type": "text",
                "validation": {"format": "year", "maxLength": 4}}
    if key == "promedio_final":
        return {"key": key, "type": "text",
                "validation": {"format": "decimal", "maxLength": 6}}
    if key.startswith("scale_"):
        return {"key": key, "type": "scale", "scale": {"min": 1, "max": 5}}
    if key == "comentario_sugerencia":
        return {"key": key, "type": "textarea", "validation": {"maxLength": 2000}}
    if key in _TEXT:
        return {"key": key, "type": "text", "validation": {"maxLength": _TEXT[key]}}
    opciones = _OPCIONES.get(key, ["Si", "No"])
    return {"key": key, "type": "radio",
            "options": [{"value": o, "label": o} for o in opciones]}


SCHEMA = {"enabled": True, "sections": [],
          "fields": [_field(k) for k in KEYS
                     if k is not None and k != "extra_aspecto_no_trabajo"]}

ORANGE = "FFFFC000"


def fila(ms_id: int, *, control="99600001", completed=None, nombre="EGRESADO SINTETICO",
         answers: dict | None = None, orange=False, theme_accent4=False) -> dict:
    """Una fila sintética. `answers` por `field_key`; lo que falte va vacío."""
    return {"id": ms_id, "control": control,
            "completed": completed or datetime(2026, 6, 15, 10, 30, 0),
            "nombre": nombre, "answers": answers or {}, "orange": orange,
            "theme_accent4": theme_accent4}


def build_xlsx(filas: list[dict], *, headers: list[str] | None = None,
               sheet: str = "Sheet1", otras_hojas=("carta de liberacion",)) -> bytes:
    """`.xlsx` en memoria con `openpyxl`, misma forma que el de Forms."""
    import openpyxl
    from openpyxl.styles import PatternFill
    from openpyxl.styles.colors import Color

    headers = headers or HEADERS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for extra in otras_hojas:
        wb.create_sheet(extra)
    ws.append(headers)
    for f in filas:
        valores = [f["id"], f["completed"], f["completed"], "anonymous", None,
                   f["nombre"], f["control"]]
        for key in KEYS[7:]:
            valores.append(f["answers"].get(key))
        ws.append(valores)
        celda = ws.cell(row=ws.max_row, column=1)
        if f["orange"]:
            celda.fill = PatternFill(fill_type="solid", fgColor=ORANGE, bgColor=ORANGE)
        elif f["theme_accent4"]:
            celda.fill = PatternFill(fill_type="solid", fgColor=Color(theme=7))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_egresados_form(make_survey_form):
    """Formulario `egresados` ABIERTO con el esquema sintético (cierra, dentro
    de la transacción de la prueba, el que haya abierto en la base de dev)."""
    import uuid
    return make_survey_form(code="egresados", version=100000 + uuid.uuid4().int % 800000,
                            status="open", schema=SCHEMA)
