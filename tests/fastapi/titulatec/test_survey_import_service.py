"""`SurveyImportService` (Tarea 2, spec `2026-10-05-titulatec-import-encuesta-
xlsx-design.md` §4.2/§4.3): lee el Excel de Microsoft Forms, liga cada columna
a su pregunta de la encuesta `egresados`, normaliza, deduplica, guarda la
respuesta como `identity_source='import'` y libera por la maquinaria de
constancias previas (`PriorClearanceService`), ligando la respuesta y la marca
de constancia por recoger.

Todo con un `.xlsx` SINTÉTICO (`_survey_xlsx.py`): el archivo real nunca.
"""
from __future__ import annotations

import re
from datetime import datetime, timedelta
from unittest.mock import patch

import pytest

import itcj2.models  # noqa: F401
from tests.fastapi.titulatec._survey_xlsx import (
    HEADERS, KEYS, build_xlsx, fila, make_egresados_form,
)

NOTIFY = "itcj2.apps.titulatec.services.notify.notify_student"
HOY_FIJO = datetime(2026, 10, 1, 10, 0, 0)
_MODULOS_RELOJ = (
    "itcj2.apps.titulatec.services.library_clearance_service",
    "itcj2.apps.titulatec.services.survey_review_service",
    "itcj2.apps.titulatec.services.prior_clearance_service",
)


def _svc():
    from itcj2.apps.titulatec.services.survey_import_service import SurveyImportService
    return SurveyImportService


def _review_svc():
    from itcj2.apps.titulatec.services.survey_review_service import SurveyReviewService
    return SurveyReviewService


def _prior_svc():
    from itcj2.apps.titulatec.services.prior_clearance_service import PriorClearanceService
    return PriorClearanceService


@pytest.fixture()
def reloj(monkeypatch):
    for modulo in _MODULOS_RELOJ:
        monkeypatch.setattr(f"{modulo}.db_now", lambda: HOY_FIJO)
    return HOY_FIJO.date()


@pytest.fixture()
def form(make_survey_form):
    return make_egresados_form(make_survey_form)


@pytest.fixture()
def proceso(make_student, make_process, make_cohort):
    def _make(*, control_number=None, current_phase=1):
        student = make_student(control_number=control_number)
        return make_process(student, cohort=make_cohort(), current_phase=current_phase)
    return _make


def _importar(db, filas, *, dry_run=False, source="egresados.xlsx"):
    rows = _svc().read_xlsx(build_xlsx(filas))
    with patch(NOTIFY):
        return _svc().import_rows(db, rows, source=source, dry_run=dry_run)


def _controles(out, bote):
    return [r["control_number"] for r in out[bote]]


def _respuestas(db, form):
    from itcj2.apps.titulatec.models import SurveyResponse
    return (db.query(SurveyResponse).filter_by(form_id=form.id)
            .order_by(SurveyResponse.id).all())


def _certs(db, review_id):
    from itcj2.apps.titulatec.models import Certificate
    return (db.query(Certificate)
            .filter_by(source_ref=f"survey_review:{review_id}")
            .order_by(Certificate.id).all())


def _certs_de_proceso(db, process_id):
    from itcj2.apps.titulatec.models import Certificate
    return db.query(Certificate).filter_by(process_id=process_id).all()


def _answer(db, response_id, key):
    from itcj2.apps.titulatec.models import SurveyAnswer
    return db.query(SurveyAnswer).filter_by(response_id=response_id, field_key=key).one()


# ---------------------------------------------------------------------------
# Lectura
# ---------------------------------------------------------------------------
class TestLectura:
    def test_mapeo_completo_de_los_69_encabezados(self):
        respuestas = {k: f"v-{k}" for k in KEYS[7:]}
        rows = _svc().read_xlsx(build_xlsx([fila(7, control="99600007",
                                                 answers=respuestas)]))
        assert len(rows) == 1
        row = rows[0]
        assert row["ms_id"] == 7
        assert row["control_raw"] == "99600007"
        assert row["full_name"] == "EGRESADO SINTETICO"
        assert row["completed_at"] == datetime(2026, 6, 15, 10, 30, 0)
        assert row["paper_pending"] is False
        esperadas = {k for k in KEYS if k is not None}
        assert set(row["answers_raw"]) == esperadas
        assert len(esperadas) == 64
        for k in KEYS[7:]:
            assert row["answers_raw"][k] == f"v-{k}"
        assert row["answers_raw"]["no_control"] == "99600007"
        assert row["answers_raw"]["nombre_completo"] == "EGRESADO SINTETICO"

    def test_encabezado_desconocido_aborta(self):
        headers = list(HEADERS)
        headers[20] = "Pregunta que no existe en la encuesta"
        with pytest.raises(ValueError, match="Pregunta que no existe"):
            _svc().read_xlsx(build_xlsx([fila(1)], headers=headers))

    def test_encabezado_faltante_aborta(self):
        headers = HEADERS[:8] + HEADERS[9:] + [None]     # sin «Sexo»
        with pytest.raises(ValueError, match="falta.*sexo"):
            _svc().read_xlsx(build_xlsx([fila(1)], headers=headers))

    def test_encabezado_repetido_aborta(self):
        headers = list(HEADERS) + [HEADERS[8]]           # «Sexo» dos veces
        with pytest.raises(ValueError, match="repetido.*Sexo") as exc:
            _svc().read_xlsx(build_xlsx([fila(1)], headers=headers))
        assert "falta" not in str(exc.value)

    def test_naranja_se_lee_en_la_columna_del_id_por_encabezado(self):
        """Si Forms mueve la columna `Id`, el naranja se busca donde quedó
        `Id` (hallado por encabezado), no en la letra A."""
        import io

        import openpyxl
        from openpyxl.styles import PatternFill

        wb = openpyxl.load_workbook(io.BytesIO(build_xlsx([
            fila(1, control="99600001", orange=True), fila(2, control="99600002")])))
        ws = wb["Sheet1"]
        ws.insert_cols(1)
        ws.cell(row=1, column=1, value="Start time")     # ignorada, puede repetirse
        ws.cell(row=3, column=1).fill = PatternFill(fill_type="solid", fgColor="FFFFC000")
        buf = io.BytesIO()
        wb.save(buf)
        rows = _svc().read_xlsx(buf.getvalue())
        assert [(r["ms_id"], r["paper_pending"]) for r in rows] == [(1, True), (2, False)]

    def test_completion_time_como_texto(self):
        rows = _svc().read_xlsx(build_xlsx([
            fila(1, completed="15/06/2026 10:30:00"),
            fila(2, completed="2026-06-16 08:00:00"),
            fila(3, completed="ayer por la tarde")]))
        assert rows[0]["completed_at"] == datetime(2026, 6, 15, 10, 30, 0)
        assert rows[1]["completed_at"] == datetime(2026, 6, 16, 8, 0, 0)
        assert rows[2]["completed_at"] is None
        assert rows[2]["completed_raw"] == "ayer por la tarde"

    def test_relleno_naranja_en_a_marca_papel_por_recoger(self):
        rows = _svc().read_xlsx(build_xlsx([
            fila(1, control="99600001", orange=True),
            fila(2, control="99600002"),
            fila(3, control="99600003", theme_accent4=True),
        ]))
        assert [r["paper_pending"] for r in rows] == [True, False, True]

    def test_hoja_inexistente_lista_las_hojas(self):
        with pytest.raises(ValueError, match="carta de liberacion"):
            _svc().read_xlsx(build_xlsx([fila(1)]), sheet="Otra")

    def test_filas_vacias_se_ignoran(self):
        import io

        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(build_xlsx([fila(1)])))
        wb["Sheet1"].append([None] * 69)
        buf = io.BytesIO()
        wb.save(buf)
        assert len(_svc().read_xlsx(buf.getvalue())) == 1


# ---------------------------------------------------------------------------
# Normalización por tipo
# ---------------------------------------------------------------------------
_RADIO = {"key": "recibir_correos", "type": "radio",
          "options": [{"value": "Si", "label": "Si"}, {"value": "No", "label": "No"}]}
_RADIO_EC = {"key": "estado_civil", "type": "radio",
             "options": [{"value": "Unión libre", "label": "Unión libre"}]}


class TestNormalizacion:
    @pytest.mark.parametrize("crudo, esperado", [
        ("Sí", "Si"), ("si", "Si"), ("  SI ", "Si"), ("no", "No"),
    ])
    def test_radio_sin_acentos_ni_mayusculas(self, crudo, esperado):
        assert _svc().normalize(_RADIO, crudo) == (esperado, False)

    def test_radio_con_acento_en_la_opcion(self):
        assert _svc().normalize(_RADIO_EC, "union  LIBRE") == ("Unión libre", False)

    def test_radio_ignora_puntuacion(self):
        campo = {"key": "periodo_egreso", "type": "radio",
                 "options": [{"value": "AGOSTO DICIEMBRE", "label": "AGOSTO DICIEMBRE"}]}
        assert _svc().normalize(campo, "Agosto-Diciembre") == ("AGOSTO DICIEMBRE", False)
        assert _svc().normalize(campo, "agosto / diciembre.") == ("AGOSTO DICIEMBRE", False)

    def test_radio_espacio_doble(self):
        campo = {"key": "sector_empresa", "type": "radio",
                 "options": [{"value": "Terciario (Educación)",
                              "label": "Terciario (Educación)"}]}
        assert _svc().normalize(campo, "Terciario  (Educación)") == (
            "Terciario (Educación)", False)

    @pytest.mark.parametrize("opcion, crudo", [
        ("Aprobé nivel III", "Aprobó nivel III"),
        ("Aprobó examen TOEFL", "Aprobé examen TOEFL"),
    ])
    def test_sinonimo_aprobo_aprobe(self, opcion, crudo):
        campo = {"key": "acreditacion_idioma", "type": "radio",
                 "options": [{"value": opcion, "label": opcion}]}
        assert _svc().normalize(campo, crudo) == (opcion, False)

    def test_schema_mal_formado_no_lanza(self):
        campo = {"key": "scale_titulado", "type": "scale",
                 "scale": {"min": "uno", "max": None}}
        assert _svc().normalize(campo, 3) == ("3", True)
        texto = {"key": "x", "type": "text", "validation": {"maxLength": "mucho"}}
        assert _svc().normalize(texto, "hola") == ("hola", True)
        assert _svc().normalize(None, "hola") == ("hola", False)

    def test_radio_sin_match_se_guarda_crudo(self):
        assert _svc().normalize(_RADIO, "No trabajo") == ("No trabajo", True)

    def test_vacio_es_none(self):
        assert _svc().normalize(_RADIO, "   ") == (None, False)
        assert _svc().normalize(_RADIO, None) == (None, False)

    def test_fecha_datetime_y_texto(self):
        campo = {"key": "fecha_nacimiento", "type": "date"}
        assert _svc().normalize(campo, datetime(1999, 5, 4)) == ("1999-05-04", False)
        assert _svc().normalize(campo, "4/5/1999") == ("1999-05-04", False)
        assert _svc().normalize(campo, "1999-05-04") == ("1999-05-04", False)
        assert _svc().normalize(campo, "ayer") == ("ayer", True)

    def test_scale(self):
        campo = {"key": "scale_titulado", "type": "scale", "scale": {"min": 1, "max": 5}}
        assert _svc().normalize(campo, 4) == (4, False)
        assert _svc().normalize(campo, "5") == (5, False)
        assert _svc().normalize(campo, 3.0) == (3, False)
        assert _svc().normalize(campo, 7) == ("7", True)
        assert _svc().normalize(campo, "Mucho 5") == (5, False)
        assert _svc().normalize(campo, "Poco 1") == (1, False)
        assert _svc().normalize(campo, " 2 ") == (2, False)
        assert _svc().normalize(campo, "1 a 5") == ("1 a 5", True)
        assert _svc().normalize(campo, "mucho") == ("mucho", True)

    def test_texto_recorta_y_largo_excedido_es_raw(self):
        campo = {"key": "nombre_empresa", "type": "text", "validation": {"maxLength": 5}}
        assert _svc().normalize(campo, "  ACME ") == ("ACME", False)
        assert _svc().normalize(campo, "ACME SA DE CV") == ("ACME SA DE CV", True)

    def test_numeros(self):
        anio = {"key": "anio_egreso", "type": "text",
                "validation": {"format": "year", "maxLength": 4}}
        prom = {"key": "promedio_final", "type": "text",
                "validation": {"format": "decimal", "maxLength": 6}}
        assert _svc().normalize(anio, 2020) == ("2020", False)
        assert _svc().normalize(anio, "2020") == ("2020", False)
        assert _svc().normalize(anio, "dos mil") == ("dos mil", True)
        assert _svc().normalize(prom, 93) == ("93", False)
        assert _svc().normalize(prom, 93.0) == ("93", False)
        assert _svc().normalize(prom, 93.5) == ("93.5", False)
        assert _svc().normalize(prom, "93,5") == ("93.5", False)
        assert _svc().normalize(prom, "noventa") == ("noventa", True)

    def test_no_control(self):
        campo = {"key": "no_control", "type": "text"}
        assert _svc().normalize(campo, " l12345678 ") == ("L12345678", False)
        assert _svc().normalize(campo, 20111222) == ("20111222", False)
        assert _svc().normalize(campo, "1234567") == ("1234567", True)


# ---------------------------------------------------------------------------
# Visibilidad (visible_when) por fila
# ---------------------------------------------------------------------------
_CAMPOS_VIS = {
    "actividad_actual": {"key": "actividad_actual", "type": "radio", "options": [
        {"value": v, "label": v} for v in ("Estudia", "Trabaja", "Estudia y trabaja",
                                           "No estudia, ni trabaja")]},
    "tipo_estudio": {"key": "tipo_estudio", "type": "radio",
                     "options": [{"value": "Maestría", "label": "Maestría"}],
                     "visible_when": {"actividad_actual": ["Estudia", "Estudia y trabaja"]}},
    "nombre_empresa": {"key": "nombre_empresa", "type": "text",
                       "validation": {"maxLength": 200},
                       "visible_when": {"actividad_actual": ["Trabaja", "Estudia y trabaja"]}},
    "scale_titulado": {"key": "scale_titulado", "type": "scale",
                       "scale": {"min": 1, "max": 5},
                       "visible_when": {"actividad_actual": ["Trabaja", "Estudia y trabaja"]}},
}


class TestVisibilidad:
    def _n(self, answers):
        return _svc().normalize_answers(_CAMPOS_VIS, answers)

    def test_radio_oculto_con_centinela_no_se_guarda(self):
        out = self._n({"actividad_actual": "Trabaja", "tipo_estudio": "No estudio"})
        assert "tipo_estudio" not in out
        assert out["actividad_actual"] == ("Trabaja", False)

    def test_radio_oculto_con_opcion_valida_se_conserva_normalizado(self):
        # Ruling de la revisión final (D1): un valor REAL en un campo oculto
        # no se pierde; `is_raw` solo marca lo que no coincide con el campo.
        stats = {}
        out = _svc().normalize_answers(
            _CAMPOS_VIS, {"actividad_actual": "Trabaja", "tipo_estudio": "Maestría"},
            stats=stats)
        assert out["tipo_estudio"] == ("Maestría", False)
        assert stats == {"cells": 2, "raw": 0, "hidden_kept": 1}

    def test_radio_oculto_que_no_coincide_queda_raw(self):
        stats = {}
        out = _svc().normalize_answers(
            _CAMPOS_VIS, {"actividad_actual": "Trabaja", "tipo_estudio": "Doctorado"},
            stats=stats)
        assert out["tipo_estudio"] == ("Doctorado", True)
        assert stats == {"cells": 2, "raw": 1, "hidden_kept": 1}

    def test_radio_visible_con_centinela_queda_raw(self):
        out = self._n({"actividad_actual": "Estudia", "tipo_estudio": "Ninguno"})
        assert out["tipo_estudio"] == ("Ninguno", True)

    @pytest.mark.parametrize("centinela", ["No trabajo", "no trabajo.", "N0", "NO"])
    def test_texto_oculto_centinela_no_se_guarda(self, centinela):
        stats = {}
        out = _svc().normalize_answers(
            _CAMPOS_VIS, {"actividad_actual": "Estudia", "nombre_empresa": centinela},
            stats=stats)
        assert "nombre_empresa" not in out
        assert stats["hidden_kept"] == 0

    def test_texto_oculto_con_dato_se_conserva_sin_raw(self):
        out = self._n({"actividad_actual": "Estudia", "nombre_empresa": "ACME"})
        assert out["nombre_empresa"] == ("ACME", False)

    def test_escala_oculta_valida_se_normaliza_sin_raw(self):
        out = self._n({"actividad_actual": "Estudia", "scale_titulado": "Mucho 5"})
        assert out["scale_titulado"] == (5, False)

    @pytest.mark.parametrize("centinela", ["No estudio", "Desempleado (a)", "Ninguno"])
    def test_radio_oculto_con_centinela_tampoco_se_guarda(self, centinela):
        out = self._n({"actividad_actual": "Trabaja", "tipo_estudio": centinela})
        assert "tipo_estudio" not in out

    def test_visible_no_cuenta_como_oculta(self):
        stats = {}
        _svc().normalize_answers(
            _CAMPOS_VIS, {"actividad_actual": "Estudia", "tipo_estudio": "Maestría"},
            stats=stats)
        assert stats == {"cells": 2, "raw": 0, "hidden_kept": 0}

    def test_fuente_raw_no_oculta_nada(self):
        out = self._n({"actividad_actual": "Jubilado", "tipo_estudio": "No estudio"})
        assert out["actividad_actual"] == ("Jubilado", True)
        assert out["tipo_estudio"] == ("No estudio", True)


def _form_con_oculta(make_survey_form):
    """`egresados` abierto donde `nombre_empresa` solo se ve si trabaja."""
    import copy
    import uuid

    from tests.fastapi.titulatec._survey_xlsx import SCHEMA

    schema = copy.deepcopy(SCHEMA)
    for campo in schema["fields"]:
        if campo["key"] == "nombre_empresa":
            campo["visible_when"] = {"actividad_actual": ["Trabaja", "Estudia y trabaja"]}
    return make_survey_form(code="egresados", version=100000 + uuid.uuid4().int % 800000,
                            status="open", schema=schema)


class TestOcultasConValorReal:
    """Ruling de la revisión final (D1): una oculta con valor REAL se guarda
    normalizada (raw solo si no coincide); una oculta centinela no. El dry-run
    las cuenta igual que la real."""

    def test_dry_run_y_real_cuentan_y_guardan_sin_raw(self, db_session, reloj,
                                                  make_survey_form):
        form = _form_con_oculta(make_survey_form)
        filas = [fila(91, control="99600131", answers={
                     "actividad_actual": "Estudia", "nombre_empresa": "ACME"}),
                 fila(92, control="99600132", answers={
                     "actividad_actual": "Estudia", "nombre_empresa": "No trabajo"})]
        rows = _svc().read_xlsx(build_xlsx(filas))

        en_seco: dict = {}
        with patch(NOTIFY):
            _svc().import_rows(db_session, rows, source="e.xlsx", dry_run=True,
                               stats=en_seco)
        assert en_seco["hidden_kept"] == 1
        assert _respuestas(db_session, form) == []

        real: dict = {}
        with patch(NOTIFY):
            _svc().import_rows(db_session, rows, source="e.xlsx", stats=real)
        assert real == en_seco

        por_control = {r.control_number: r for r in _respuestas(db_session, form)}
        acme = _answer(db_session, por_control["99600131"].id, "nombre_empresa")
        assert acme.is_raw is False and acme.value_text == "ACME"
        assert por_control["99600131"].answers["nombre_empresa"] == "ACME"
        assert "nombre_empresa" not in por_control["99600132"].answers


# ---------------------------------------------------------------------------
# Importación: escritura y liberación
# ---------------------------------------------------------------------------
class TestImportacion:
    def test_con_proceso_libera_liga_respuesta_y_papel(
            self, db_session, reloj, form, proceso):
        proc = proceso(control_number="99600101")
        out = _importar(db_session, [fila(11, control="99600101", orange=True, answers={
            "sexo": "hombre", "recibir_correos": "No trabajo", "scale_titulado": 4,
            "fecha_nacimiento": datetime(1999, 5, 4),
            "extra_aspecto_no_trabajo": "3"})])

        assert _controles(out, "released") == ["99600101"]
        (resp,) = _respuestas(db_session, form)
        assert resp.identity_source == "import"
        assert resp.import_ref == "msforms:11:2026-06-15T10:30:00"
        assert resp.control_number == "99600101"
        assert resp.submitted_at == datetime(2026, 6, 15, 10, 30, 0)
        assert resp.form_version == form.version
        assert resp.user_id == proc.student_id
        assert resp.process_id == proc.id
        assert resp.cohort_id == proc.cohort_id
        assert resp.answers["sexo"] == "Hombre"
        assert resp.answers["nombre_completo"] == "EGRESADO SINTETICO"
        assert resp.answers["extra_aspecto_no_trabajo"] == "3"
        assert _answer(db_session, resp.id, "recibir_correos").is_raw is True
        assert _answer(db_session, resp.id, "sexo").is_raw is False
        assert float(_answer(db_session, resp.id, "scale_titulado").value_num) == 4
        assert _answer(db_session, resp.id, "fecha_nacimiento").value_text == "1999-05-04"

        review = _review_svc().get_for_process(db_session, proc.id)
        assert review.origin == "prior"
        assert review.response_id == resp.id
        assert review.paper_pending is True
        assert review.prior_issued_on == datetime(2026, 6, 15).date()
        # Spec folios 2026-10-05: la liberada también deja su folio (GTV del
        # semestre anterior al del registro, sin emisor: es una importación).
        (cert,) = _certs(db_session, review.id)
        assert re.fullmatch(r"GTV-2026A-\d{4}", cert.number), cert.number
        assert cert.issued_by_id is None
        assert cert.control_number == "99600101"

    def test_sin_proceso_difiere_y_apply_pending_liga_al_inscribirse(
            self, db_session, reloj, form, proceso):
        from itcj2.apps.titulatec.models import PriorClearance

        out = _importar(db_session, [fila(12, control="99600102", orange=True)])
        assert _controles(out, "deferred") == ["99600102"]
        (resp,) = _respuestas(db_session, form)
        assert resp.process_id is None and resp.user_id is None
        previa = (db_session.query(PriorClearance)
                  .filter_by(kind="survey", control_number="99600102").one())
        assert previa.response_id == resp.id
        assert previa.paper_pending is True
        assert previa.issued_on == datetime(2026, 6, 15).date()
        # `created_at` es el `NOW()` REAL de Postgres (el `reloj` solo parchea
        # `db_now()`); el folio de una diferida sale de esa fecha (D5 de los
        # folios), así que se fija al día de la prueba: sin esto el semestre
        # esperado (2026A) cambiaría con el reloj real a partir de 2027.
        previa.created_at = HOY_FIJO
        db_session.flush()

        proc = proceso(control_number="99600102")
        with patch(NOTIFY):
            assert _prior_svc().apply_pending(db_session, proc, "99600102") == ["survey"]
        review = _review_svc().get_for_process(db_session, proc.id)
        assert review.response_id == resp.id
        assert review.paper_pending is True
        (cert,) = _certs(db_session, review.id)           # la diferida también folia
        assert cert.number.startswith("GTV-2026A-")
        db_session.refresh(resp)
        assert resp.user_id == proc.student_id
        assert resp.process_id == proc.id
        assert resp.cohort_id == proc.cohort_id

    def test_usuario_sin_proceso_queda_con_user_id(
            self, db_session, reloj, form, make_student):
        alumno = make_student(control_number="99600103")
        out = _importar(db_session, [fila(13, control="99600103")])
        assert _controles(out, "deferred") == ["99600103"]
        (resp,) = _respuestas(db_session, form)
        assert resp.user_id == alumno.id
        assert resp.process_id is None

    def test_duplicados_se_queda_la_mas_reciente(self, db_session, reloj, form):
        out = _importar(db_session, [
            fila(21, control="99600104", completed=datetime(2026, 6, 1, 9, 0)),
            fila(22, control="99600104", completed=datetime(2026, 6, 20, 9, 0)),
            fila(23, control="99600104", completed=datetime(2026, 6, 10, 9, 0)),
        ])
        (resp,) = _respuestas(db_session, form)
        assert resp.import_ref == "msforms:22:2026-06-20T09:00:00"
        assert sorted(r["ms_id"] for r in out["duplicates"]) == [21, 23]
        assert _controles(out, "deferred") == ["99600104"]

    def test_idempotente_segunda_corrida_ya_importadas(self, db_session, reloj, form):
        from itcj2.apps.titulatec.models import PriorClearance, SurveyAnswer

        filas = [fila(31, control="99600105"), fila(32, control="")]
        _importar(db_session, filas)
        n_answers = db_session.query(SurveyAnswer).count()
        out = _importar(db_session, filas)
        assert len(_respuestas(db_session, form)) == 2
        assert db_session.query(SurveyAnswer).count() == n_answers
        assert sorted(r["ms_id"] for r in out["already_imported"]) == [31, 32]
        for bote in ("released", "deferred", "saved_unreleased"):
            assert out[bote] == []
        assert (db_session.query(PriorClearance)
                .filter_by(kind="survey", control_number="99600105").count()) == 1

    @pytest.mark.parametrize("control", ["", "1234567", "ABC"])
    def test_control_invalido_guarda_sin_liberar(self, db_session, reloj, form, control):
        from itcj2.apps.titulatec.models import PriorClearance

        n_previas = db_session.query(PriorClearance).count()
        out = _importar(db_session, [fila(41, control=control)])
        assert len(out["saved_unreleased"]) == 1
        (resp,) = _respuestas(db_session, form)
        assert resp.identity_source == "import"
        assert resp.control_number is None
        assert db_session.query(PriorClearance).count() == n_previas

    def test_vencida_guarda_sin_liberar(self, db_session, reloj, form, proceso):
        proc = proceso(control_number="99600106")
        out = _importar(db_session, [fila(42, control="99600106",
                                          completed=datetime(2025, 6, 1, 9, 0))])
        assert _controles(out, "saved_unreleased") == ["99600106"]
        assert len(_respuestas(db_session, form)) == 1
        assert _review_svc().get_for_process(db_session, proc.id) is None

    def test_conflicto_guarda_y_no_pisa_la_revision(
            self, db_session, reloj, proceso, make_survey_review, make_survey_form):
        proc = proceso(control_number="99600107")
        review = make_survey_review(proc, status="in_review")
        original = review.response_id
        # `make_survey_review` abre su propio `egresados`: el del import va DESPUÉS.
        form = make_egresados_form(make_survey_form)
        out = _importar(db_session, [fila(51, control="99600107", orange=True)])
        assert _controles(out, "conflicts") == ["99600107"]
        (nueva,) = _respuestas(db_session, form)
        assert nueva.id != original
        # Ruling: user_id y cohort_id aunque no se ligue; process_id no.
        assert nueva.user_id == proc.student_id
        assert nueva.cohort_id == proc.cohort_id
        assert nueva.process_id is None
        db_session.refresh(review)
        assert review.status == "in_review"
        assert review.response_id == original
        assert review.paper_pending is False

    def test_ya_liberada_por_csv_adjunta_la_respuesta(
            self, db_session, reloj, form, proceso):
        proc = proceso(control_number="99600108")
        with patch(NOTIFY):
            _prior_svc().import_rows(db_session, kind="survey", source="viejo.csv", rows=[
                {"control_number": "99600108", "issued_on": "2026-05-01"}])
        review = _review_svc().get_for_process(db_session, proc.id)
        assert review.response_id is None

        out = _importar(db_session, [fila(61, control="99600108", orange=True)])
        assert _controles(out, "already_released") == ["99600108"]
        (resp,) = _respuestas(db_session, form)
        db_session.refresh(review)
        assert review.response_id == resp.id
        assert review.paper_pending is True
        assert resp.process_id == proc.id

    def test_ya_liberada_con_respuesta_propia_no_se_toca(
            self, db_session, reloj, proceso, make_survey_review, make_survey_form):
        proc = proceso(control_number="99600109")
        review = make_survey_review(proc, status="approved")
        original = review.response_id
        form = make_egresados_form(make_survey_form)
        out = _importar(db_session, [fila(62, control="99600109", orange=True)])
        assert _controles(out, "already_released") == ["99600109"]
        db_session.refresh(review)
        assert review.response_id == original
        assert review.paper_pending is False
        (resp,) = _respuestas(db_session, form)
        assert resp.process_id is None
        assert resp.cohort_id == proc.cohort_id

    def test_ya_aplicada_en_proceso_cerrado_adjunta_a_la_previa_y_su_revision(
            self, db_session, reloj, form, proceso):
        from itcj2.apps.titulatec.models import PriorClearance, ProcessPhase

        # Una diferida del CSV que ya se aplicó; luego la fase 2 se aprobó
        # (el proceso deja de estar «abierto» para la importación).
        with patch(NOTIFY):
            _prior_svc().import_rows(db_session, kind="survey", source="viejo.csv", rows=[
                {"control_number": "99600110", "issued_on": "2026-06-20"}])
        proc = proceso(control_number="99600110")
        with patch(NOTIFY):
            _prior_svc().apply_pending(db_session, proc, "99600110")
        (db_session.query(ProcessPhase).filter_by(process_id=proc.id, phase_number=2)
         .update({"status": "approved"}))
        db_session.flush()

        out = _importar(db_session, [fila(63, control="99600110")])  # 2026-06-15
        assert _controles(out, "already_released") == ["99600110"]
        (resp,) = _respuestas(db_session, form)
        review = _review_svc().get_for_process(db_session, proc.id)
        assert review.response_id == resp.id
        previa = (db_session.query(PriorClearance)
                  .filter_by(kind="survey", control_number="99600110").one())
        assert previa.response_id == resp.id
        db_session.refresh(resp)
        assert resp.process_id == proc.id

    def test_dry_run_no_escribe_nada(self, db_session, reloj, form, proceso):
        from itcj2.apps.titulatec.models import PriorClearance, SurveyAnswer, SurveyReview

        proc = proceso(control_number="99600111")
        cuenta = lambda: (len(_respuestas(db_session, form)),  # noqa: E731
                          db_session.query(SurveyAnswer).count(),
                          db_session.query(PriorClearance).count(),
                          db_session.query(SurveyReview).count())
        antes = cuenta()
        filas = [fila(71, control="99600111"), fila(72, control="99600112"),
                 fila(73, control="99600112", completed=datetime(2026, 5, 1)),
                 fila(74, control="")]
        out = _importar(db_session, filas, dry_run=True)
        assert cuenta() == antes
        assert _controles(out, "released") == ["99600111"]
        assert _controles(out, "deferred") == ["99600112"]
        assert len(out["duplicates"]) == 1
        assert len(out["saved_unreleased"]) == 1
        assert _review_svc().get_for_process(db_session, proc.id) is None

    def test_sin_formulario_abierto_falla(self, db_session, reloj, make_survey_form):
        from itcj2.apps.titulatec.models import SurveyForm
        (db_session.query(SurveyForm).filter_by(code="egresados", status="open")
         .update({"status": "closed"}))
        db_session.flush()
        rows = _svc().read_xlsx(build_xlsx([fila(1)]))
        with pytest.raises(ValueError, match="egresados"):
            _svc().import_rows(db_session, rows, source="x.xlsx", dry_run=True)

    def test_completion_time_ilegible_es_invalida_y_no_se_guarda(
            self, db_session, reloj, form):
        out = _importar(db_session, [fila(91, control="99600130", completed="ayer"),
                                     fila(92, control="99600131",
                                          completed="15/06/2026 10:30:00")])
        assert [r["ms_id"] for r in out["invalid"]] == [91]
        assert _controles(out, "deferred") == ["99600131"]
        (resp,) = _respuestas(db_session, form)
        assert resp.import_ref == "msforms:92:2026-06-15T10:30:00"

    def test_idempotencia_entre_versiones_del_formulario(
            self, db_session, reloj, form, make_survey_form):
        _importar(db_session, [fila(93, control="99600132")])
        nueva_version = make_egresados_form(make_survey_form)
        out = _importar(db_session, [fila(93, control="99600132")])
        assert [r["ms_id"] for r in out["already_imported"]] == [93]
        assert _respuestas(db_session, nueva_version) == []

    def test_previa_pendiente_conserva_fecha_mas_nueva_y_nota(
            self, db_session, reloj, form):
        from itcj2.apps.titulatec.models import PriorClearance

        _prior_svc().import_rows(db_session, kind="survey", source="viejo.csv", rows=[
            {"control_number": "99600133", "issued_on": "2026-07-01",
             "note": "Nota del CSV"}])
        out = _importar(db_session, [fila(94, control="99600133")])  # 2026-06-15
        assert _controles(out, "deferred") == ["99600133"]
        previa = (db_session.query(PriorClearance)
                  .filter_by(kind="survey", control_number="99600133").one())
        assert previa.issued_on == datetime(2026, 7, 1).date()
        assert previa.note == "Nota del CSV"
        (resp,) = _respuestas(db_session, form)
        assert previa.response_id == resp.id

    def test_previa_pendiente_toma_la_fecha_si_es_mas_nueva(self, db_session, reloj, form):
        from itcj2.apps.titulatec.models import PriorClearance

        _prior_svc().import_rows(db_session, kind="survey", source="viejo.csv", rows=[
            {"control_number": "99600134", "issued_on": "2026-05-01", "note": "CSV"}])
        _importar(db_session, [fila(95, control="99600134")])          # 2026-06-15
        previa = (db_session.query(PriorClearance)
                  .filter_by(kind="survey", control_number="99600134").one())
        assert previa.issued_on == datetime(2026, 6, 15).date()
        assert previa.note == "CSV"

    def test_botes_completos(self, db_session, reloj, form):
        out = _importar(db_session, [fila(81, control="99600113")], dry_run=True)
        assert tuple(out) == ("released", "deferred", "already_released", "conflicts",
                              "saved_unreleased", "duplicates", "already_imported",
                              "invalid")
        assert set(out["deferred"][0]) >= {"control_number", "reason", "ms_id"}


class TestFolios:
    """Spec folios 2026-10-05 §3.3: UN folio por fila liberada, nada por las
    demás."""

    def test_un_folio_por_fila_liberada_y_ninguno_por_las_demas(
            self, db_session, reloj, proceso, make_survey_review,
            make_survey_form):
        from itcj2.apps.titulatec.models import Certificate

        a = proceso(control_number="99600201")
        b = proceso(control_number="99600202")
        c = proceso(control_number="99600203")            # vencida: guarda sin liberar
        d = proceso(control_number="99600204")            # conflicto: GTV decide
        make_survey_review(d, status="in_review")
        make_egresados_form(make_survey_form)   # el del import va DESPUÉS del de `make_survey_review`
        antes = db_session.query(Certificate).count()

        out = _importar(db_session, [
            fila(201, control="99600201"),
            fila(202, control="99600202", orange=True),
            fila(203, control="99600203", completed=datetime(2025, 6, 1, 9, 0)),
            fila(204, control="99600204"),
            fila(205, control="99600205"),                # sin proceso: diferida
        ])

        assert _controles(out, "released") == ["99600201", "99600202"]
        assert _controles(out, "saved_unreleased") == ["99600203"]
        assert _controles(out, "conflicts") == ["99600204"]
        assert _controles(out, "deferred") == ["99600205"]
        assert db_session.query(Certificate).count() == antes + 2

        reviews = [_review_svc().get_for_process(db_session, proc.id) for proc in (a, b)]
        folios = [_certs(db_session, review.id) for review in reviews]
        assert all(len(f) == 1 for f in folios)
        numeros = [f[0].number for f in folios]
        assert all(re.fullmatch(r"GTV-2026A-\d{4}", n) for n in numeros), numeros
        assert numeros[0] != numeros[1]
        assert _certs_de_proceso(db_session, c.id) == []
        assert _certs_de_proceso(db_session, d.id) == []

    def test_dry_run_no_emite_folios(self, db_session, reloj, form, proceso):
        from itcj2.apps.titulatec.models import Certificate

        proceso(control_number="99600211")
        antes = db_session.query(Certificate).count()

        out = _importar(db_session, [fila(211, control="99600211")], dry_run=True)

        assert _controles(out, "released") == ["99600211"]
        assert db_session.query(Certificate).count() == antes

    def test_la_segunda_corrida_no_vuelve_a_emitir(self, db_session, reloj, form, proceso):
        from itcj2.apps.titulatec.models import Certificate

        proceso(control_number="99600212")
        _importar(db_session, [fila(212, control="99600212")])
        despues_de_la_primera = db_session.query(Certificate).count()

        out = _importar(db_session, [fila(212, control="99600212")])

        assert _controles(out, "already_imported") == ["99600212"]
        assert db_session.query(Certificate).count() == despues_de_la_primera


def test_prior_clearance_paper_pending_exige_true_literal(
        db_session, reloj):
    """`paper_pending` se lee con `is True`: un texto «False» no marca papel."""
    from itcj2.apps.titulatec.models import PriorClearance

    _prior_svc().import_rows(db_session, kind="survey", source="x.xlsx", rows=[
        {"control_number": "99600120", "issued_on": "2026-06-01",
         "response_id": None, "paper_pending": "False"}])
    previa = (db_session.query(PriorClearance)
              .filter_by(kind="survey", control_number="99600120").one())
    assert previa.paper_pending is False
