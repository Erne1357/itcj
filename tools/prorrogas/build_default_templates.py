"""Genera las 6 plantillas del sistema de Prórrogas a partir de las cartas reales.

Uso (dentro del contenedor backend, desde /app):

    python tools/prorrogas/build_default_templates.py \
        --source instance/tmp/prorrogas_ejemplos \
        --out itcj2/apps/prorrogas_tec/documents/defaults

- Las tres cartas compromiso salen de las tres cartas reales de Antecedentes:
  los datos del estudiante se cambian por marcadores (tabla LITERALS) y se
  agrega el lugar de la firma electrónica sobre el nombre del estudiante.
- Las otras tres (aceptación, rechazo, formato de solicitud) parten del
  membrete de la carta de 1 pago y llevan el texto de spec §6.4.

Se detiene con código 1 si un texto real no aparece tal cual o si una plantilla
generada no cumple los marcadores de su formato: nunca escribe a medias.
"""
from __future__ import annotations

import argparse
import sys
from copy import deepcopy
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from itcj2.apps.prorrogas_tec.documents.catalog import DOC_TYPES  # noqa: E402
from itcj2.apps.prorrogas_tec.services.docx_fill import (  # noqa: E402
    find_placeholders, iter_paragraphs, replace_literals,
)

BASE_NEW_LETTERS = "Prórroga un pago 2026.docx"

SOURCES = {
    "carta_compromiso_1": "Prórroga un pago 2026.docx",
    "carta_compromiso_2": "Prórroga en dos exhibiciones 2026.docx",
    "carta_compromiso_3": "Prórroga tres pagos 2026 .docx",
}

# Texto real → marcador, en orden. Las claves salen de las cartas de Antecedentes.
LITERALS = {
    "carta_compromiso_1": {
        "03/febrero/2026": "{{fecha}}",
        "$6,400.00": "{{monto_total}}",
        "enero – junio 2026": "{{semestre}}",
        "CARDENAS ALMANZA ROBERTO": "{{nombre}}",
        "M23111966": "{{no_control}}",
        "de la maestría de Ingeniería Administrativa": "de la carrera de {{carrera}}",
        "03 de febrero de 2026": "{{pago1_fecha}}",
    },
    "carta_compromiso_2": {
        "13/enero/2026": "{{fecha}}",
        "$3,150.00": "{{monto_total}}",
        "como concepto de inscripción para el estudiante":
            "como concepto de inscripción para el semestre {{semestre}}, para el (la) estudiante",
        "PAYAN FLORES ITZEL GUADALUPE": "{{nombre}}",
        "22111670": "{{no_control}}",
        "Licenciatura en Administración": "{{carrera}}",
        "$1,575.00 a efectuarse el día 26 de febrero de 2026": "{{pago1_monto}} a efectuarse el día {{pago1_fecha}}",
        "$1,575.00 a efectuarse el día 23 de abril de 2026": "{{pago2_monto}} a efectuarse el día {{pago2_fecha}}",
    },
    "carta_compromiso_3": {
        "03/febrero/2026": "{{fecha}}",
        "$3,650.00": "{{monto_total}}",
        "enero – junio 2026": "{{semestre}}",
        "HERNANDEZ SOLANO JOSE ESAU": "{{nombre}}",
        "24112010": "{{no_control}}",
        "Ingeniería en Sistemas Computacionales EAD": "{{carrera}}",
        "$1,217.00 a efectuarse el día 13 de febrero de 2026": "{{pago1_monto}} a efectuarse el día {{pago1_fecha}}",
        "$1,217.00 a efectuarse el día 27 de febrero de 2026": "{{pago2_monto}} a efectuarse el día {{pago2_fecha}}",
        "$1,216.00 a efectuarse el día 27 de marzo de 2026": "{{pago3_monto}} a efectuarse el día {{pago3_fecha}}",
    },
}

_CIERRE = ["A T E N T A M E N T E", "Excelencia en Educación Tecnológica®", "SUBDIRECCIÓN ACADÉMICA"]

# Spec §6.4, copiado exactamente.
NEW_LETTERS = {
    "carta_aceptacion": [
        "Ciudad Juárez, Chihuahua, {{fecha}}",
        "{{nombre}}",
        "No. de control {{no_control}} · {{carrera}}",
        "PRESENTE",
        "Por medio de la presente se le informa que su solicitud de prórroga de pago de inscripción del "
        "semestre {{semestre}}, registrada con el folio {{folio}}, ha sido ACEPTADA.",
        "Para continuar, acuda a la Subdirección Académica para acordar el monto y las fechas de sus pagos; "
        "con ello se emitirá su carta compromiso de pago.",
        "{{comentario}}",
        *_CIERRE,
    ],
    "carta_rechazo": [
        "Ciudad Juárez, Chihuahua, {{fecha}}",
        "{{nombre}}",
        "No. de control {{no_control}} · {{carrera}}",
        "PRESENTE",
        "Por medio de la presente se le informa que su solicitud de prórroga de pago de inscripción del "
        "semestre {{semestre}}, registrada con el folio {{folio}}, no fue autorizada por el siguiente motivo:",
        "{{motivo_rechazo}}",
        "Si su situación cambia, puede acudir nuevamente a la Subdirección Académica.",
        *_CIERRE,
    ],
    "formato_solicitud": [
        "Ciudad Juárez, Chihuahua, {{fecha_solicitud}}",
        "JEOVANNY RAFAEL RODRÍGUEZ MEJÍA",
        "SUBDIRECTOR ACADÉMICO",
        "PRESENTE",
        "Quien suscribe, {{nombre}}, con número de control {{no_control}}, estudiante de {{carrera}}, "
        "solicito de la manera más atenta una prórroga para cubrir el pago de inscripción del semestre "
        "{{semestre}} en {{plazos_solicitados}} pago(s), por el siguiente motivo:",
        "{{motivo}}",
        "Correo de contacto: {{correo}}",
        "A T E N T A M E N T E",
        "{{firma_estudiante}}",
        "{{nombre}}",
        "{{firma_estudiante_leyenda}}",
        "Folio {{folio}}",
    ],
}


def fail(msg: str) -> None:
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def _save(doc) -> bytes:
    for attr in ("title", "subject", "keywords", "comments"):
        setattr(doc.core_properties, attr, "")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _unique(paragraphs, text: str, code: str) -> Paragraph:
    found = [p for p in paragraphs if p.text.strip() == text]
    if len(found) != 1:
        fail(f"{code}: se esperaba exactamente un párrafo «{text}» y hay {len(found)}")
    return found[0]


def _paragraph_like(ref: Paragraph, text: str, *, size_pt: float | None = None):
    """Párrafo nuevo con el formato de párrafo y de letra de ``ref``."""
    p_el = OxmlElement("w:p")
    if ref._p.pPr is not None:
        p_el.append(deepcopy(ref._p.pPr))
    para = Paragraph(p_el, ref._parent)
    run = para.add_run(text)
    ref_runs = ref.runs
    if ref_runs and ref_runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(ref_runs[0]._r.rPr))
    if size_pt:
        run.font.size = Pt(size_pt)
    return p_el


def build_compromiso(code: str, source: bytes) -> bytes:
    data, counts = replace_literals(source, LITERALS[code])
    missing = [k for k, n in counts.items() if n == 0]
    if missing:
        fail(f"{code}: estos textos no aparecen tal cual en la carta real: {missing}")
    doc = Document(BytesIO(data))
    paragraphs = list(iter_paragraphs(doc))
    nombre = _unique(paragraphs, "{{nombre}}", code)
    estudiante = _unique(paragraphs, "ESTUDIANTE", code)
    nombre._p.addprevious(_paragraph_like(nombre, "{{firma_estudiante}}"))
    estudiante._p.addnext(_paragraph_like(estudiante, "{{firma_estudiante_leyenda}}", size_pt=7))
    return _save(doc)


def build_new_letter(code: str, base: bytes) -> bytes:
    doc = Document(BytesIO(base))
    body = doc.element.body
    paragraphs = list(iter_paragraphs(doc))
    fecha = next((p for p in paragraphs if "Ciudad Juárez, Chihuahua" in p.text), None)
    cuerpo = next((p for p in paragraphs if p.text.strip().startswith("Por medio de la presente")), None)
    if fecha is None or cuerpo is None:
        fail(f"{code}: la carta base no tiene la línea de fecha o el párrafo «Por medio de la presente»")
    if fecha._p.getparent() is not body:
        fail(f"{code}: la línea de fecha de la carta base no está en el cuerpo principal")

    # Formato de párrafo y de letra del cuerpo de la carta real.
    ref_ppr = deepcopy(cuerpo._p.pPr) if cuerpo._p.pPr is not None else None
    ref_rpr = deepcopy(cuerpo.runs[0]._r.rPr) if cuerpo.runs and cuerpo.runs[0]._r.rPr is not None else None

    children = list(body.iterchildren())
    for el in children[children.index(fecha._p):]:
        if el.tag != qn("w:sectPr"):
            body.remove(el)

    lines = NEW_LETTERS[code]
    cierre = lines.index("A T E N T A M E N T E")
    for i, line in enumerate(lines):
        if i:
            doc.add_paragraph()
        para = doc.add_paragraph()
        if ref_ppr is not None:
            para._p.insert(0, deepcopy(ref_ppr))
        run = para.add_run(line)
        if ref_rpr is not None:
            run._r.insert(0, deepcopy(ref_rpr))
        if i == 0:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif i >= cierre:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if line == "{{firma_estudiante_leyenda}}":
            run.font.size = Pt(7)
    return _save(doc)


def validate(code: str, data: bytes) -> None:
    t = DOC_TYPES[code]
    found = find_placeholders(data)
    if not t.required <= found:
        fail(f"{code}: faltan marcadores obligatorios {sorted(t.required - found)}")
    if not found <= t.allowed:
        fail(f"{code}: marcadores no permitidos {sorted(found - t.allowed)}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--source", required=True, type=Path, help="carpeta con las 3 cartas reales")
    ap.add_argument("--out", required=True, type=Path, help="carpeta de salida (documents/defaults)")
    args = ap.parse_args()

    def read(name: str) -> bytes:
        path = args.source / name
        if not path.exists():
            fail(f"no existe {path}")
        return path.read_bytes()

    outputs: dict[str, bytes] = {}
    for code, name in SOURCES.items():
        outputs[code] = build_compromiso(code, read(name))
    base = read(BASE_NEW_LETTERS)
    for code in NEW_LETTERS:
        outputs[code] = build_new_letter(code, base)
    for code, data in outputs.items():
        validate(code, data)

    args.out.mkdir(parents=True, exist_ok=True)
    for code in DOC_TYPES:
        (args.out / f"{code}.docx").write_bytes(outputs[code])
        print(f"OK {code}")


if __name__ == "__main__":
    main()
