"""Pestaña Estudiantes de la secretaria: lista por semestre con su cuenta y expediente de una solicitud."""
from __future__ import annotations

import logging
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from itcj2.apps.prorrogas_tec.documents.catalog import DOC_TYPES
from itcj2.apps.prorrogas_tec.models.event import ProrrogaEvent
from itcj2.apps.prorrogas_tec.models.plan import ProrrogaPayment
from itcj2.apps.prorrogas_tec.models.request import ProrrogaRequest
from itcj2.apps.prorrogas_tec.services import account, notify
from itcj2.apps.prorrogas_tec.services.document_service import DocumentService
from itcj2.apps.prorrogas_tec.services.errors import ServiceError
from itcj2.apps.prorrogas_tec.services.events import log_event
from itcj2.apps.prorrogas_tec.services.formatting import fecha_larga, pesos
from itcj2.apps.prorrogas_tec.services.identity import folio
from itcj2.apps.prorrogas_tec.services.periods import period_for, period_label
from itcj2.apps.prorrogas_tec.services.request_service import RequestService
from itcj2.core.models.academic_period import AcademicPeriod
from itcj2.core.utils.timezone import db_now

FILTROS = ("todos", "con_adeudo", "al_corriente", "sin_plan", "liquidada")
_SITUACION = {"con_adeudo": "atrasada", "al_corriente": "al_corriente", "sin_plan": "sin_plan",
              "liquidada": "liquidada"}
_ENTRE_RECORDATORIOS = timedelta(hours=24)
logger = logging.getLogger(__name__)
_COLUMNAS = [("No. de control", 16), ("Nombre", 36), ("Carrera", 36), ("Correo", 32),
             ("Folio de la prórroga", 20), ("Monto total", 14), ("Pagado", 14), ("Saldo", 14),
             ("Parcialidades vencidas", 22), ("Vencida desde", 16)]
_DINERO = '"$"#,##0.00'


def _tiempos(now):
    now = now or db_now()
    if isinstance(now, datetime):
        return now, now.date(), now
    return now, now, datetime.combine(now, time())


def _puede_recordar(req, c, ahora: datetime) -> bool:
    return (c.total is not None and c.saldo > 0
            and (req.last_reminded_at is None or ahora - req.last_reminded_at >= _ENTRE_RECORDATORIOS))


def _avisar_desbloqueo(req, c) -> bool:
    return req.block_exported_at is not None and c.situacion == "liquidada" and req.unblock_notified_at is None


class StudentsService:
    @staticmethod
    def periodos(db, today: date) -> dict:
        ps = (db.query(AcademicPeriod)
              .filter(AcademicPeriod.id.in_(
                  db.query(ProrrogaRequest.period_id).filter(ProrrogaRequest.status == "aceptada")))
              .order_by(AcademicPeriod.start_date.desc()).all())
        items = [{"id": p.id, "label": period_label(p)} for p in ps]
        hoy = period_for(db, today)
        ids = [i["id"] for i in items]
        actual = hoy.id if hoy and hoy.id in ids else (ids[0] if ids else None)
        return {"items": items, "actual": actual}

    @staticmethod
    def listar(db, period_id: int, *, q: str | None = None, filtro: str = "todos", page: int = 1,
               per_page: int = 25, today: date) -> dict:
        if filtro not in FILTROS:
            raise ServiceError("filtro_invalido", "Filtro no válido.", 400)
        page = max(page, 1)
        per_page = min(max(per_page, 1), 100)
        qry = db.query(ProrrogaRequest).filter(ProrrogaRequest.period_id == period_id,
                                               ProrrogaRequest.status == "aceptada")
        if q and q.strip():
            like = f"%{q.strip()}%"
            qry = qry.filter(ProrrogaRequest.control_number.ilike(like) | ProrrogaRequest.student_name.ilike(like))
        reqs = qry.all()
        # shortcut: situación calculada en Python por semestre; pasar a SQL si un semestre supera ~2000 solicitudes
        cuentas = account.load_many(db, reqs, today)
        conteos = {"todos": len(reqs)} | {k: sum(1 for c in cuentas.values() if c.situacion == s)
                                          for k, s in _SITUACION.items()}
        if filtro != "todos":
            reqs = [r for r in reqs if cuentas[r.id].situacion == _SITUACION[filtro]]
        reqs.sort(key=lambda r: ((r.student_name or "").lower(), r.id))
        pagina = reqs[(page - 1) * per_page: page * per_page]
        ahora = db_now()
        items = []
        for r in pagina:
            c = cuentas[r.id]
            items.append({
                "id": r.id, "folio": folio(r.id), "control": r.control_number, "nombre": r.student_name,
                "situacion": c.situacion, "faltantes": c.faltantes,
                "vencida_desde": c.vencida_desde.isoformat() if c.vencida_desde else None,
                "pagos": [p.estado for p in c.parcialidades],
                "pagados": sum(1 for p in c.parcialidades if p.estado == "pagada"),
                "puede_recordar": _puede_recordar(r, c, ahora),
                "avisar_desbloqueo": _avisar_desbloqueo(r, c)})
        return {"items": items, "total": len(reqs), "page": page, "per_page": per_page, "conteos": conteos}

    @staticmethod
    def expediente(db, request_id: int, today: date) -> dict:
        req = db.get(ProrrogaRequest, request_id)
        if req is None or req.status != "aceptada":
            raise ServiceError("no_encontrada", "No se encontró la solicitud.", 404)
        c = account.load(db, req, today)
        period = db.get(AcademicPeriod, req.period_id) if req.period_id else None  # histórica sin semestre: solo lectura
        conservadas = [p for p in c.parcialidades if p.estado == "pagada"]
        suma = sum((p.amount for p in conservadas), Decimal("0"))
        carta = next((d for d in DocumentService.current(db, req) if d.doc_type.startswith("carta_compromiso")), None)
        recibos = (db.query(ProrrogaPayment).filter(ProrrogaPayment.request_id == req.id)
                   .order_by(ProrrogaPayment.recorded_at.desc(), ProrrogaPayment.id.desc()).all())
        eventos = (db.query(ProrrogaEvent).filter(ProrrogaEvent.request_id == req.id)
                   .order_by(ProrrogaEvent.created_at.desc(), ProrrogaEvent.id.desc()).all())
        return {
            "solicitud": RequestService.request_out(db, req),
            "estudiante": {"nombre": req.student_name, "control": req.control_number, "correo": req.email,
                           "carrera": req.program_name},
            "cuenta": account.cuenta_json(c),
            "recibos": [{"id": p.id, "folio": p.folio, "fecha": p.paid_on.isoformat(), "monto": account.money(p.amount),
                         "registrado": p.recorded_at.isoformat(),
                         "anulado": p.voided_at.isoformat() if p.voided_at else None,
                         "motivo_anulacion": p.void_reason} for p in recibos],
            "carta": ({"doc_id": carta.id, "tipo": carta.doc_type, "version": carta.version,
                       "firmada": carta.signed_at is not None} if carta else None),
            "plan": {
                "confirmado": req.plan_confirmed_at.isoformat() if req.plan_confirmed_at else None,
                "puede_replantear": period is not None and c.situacion != "liquidada" and today <= period.end_date,
                "conservadas": [p.number for p in conservadas],
                "suma_conservadas": account.money(suma),
                "max_nuevas": 3 - len(conservadas),
                "fecha_limite": period.end_date.isoformat() if period else None},
            "bloqueo": {"exportada": req.block_exported_at.isoformat() if req.block_exported_at else None,
                        "avisar_desbloqueo": _avisar_desbloqueo(req, c),
                        "avisado": req.unblock_notified_at.isoformat() if req.unblock_notified_at else None},
            "recordatorio": {"ultimo": req.last_reminded_at.isoformat() if req.last_reminded_at else None,
                             "puede": _puede_recordar(req, c, db_now())},
            "historial": [{"tipo": e.event_type, "fecha": e.created_at.isoformat(), "datos": e.payload}
                          for e in eventos],
        }

    @staticmethod
    def _bloquear(db, request_id: int) -> ProrrogaRequest:
        req = (db.query(ProrrogaRequest).filter_by(id=request_id)
               .populate_existing().with_for_update().first())
        if req is None or req.status != "aceptada":
            db.rollback()
            raise ServiceError("no_encontrada", "No se encontró la solicitud.", status=404)
        return req

    @staticmethod
    def remind(db, request_id: int, actor_id: int, *, now: date | datetime | None = None) -> ProrrogaRequest:
        """Recordatorio de pago por correo y aviso; como máximo uno cada 24 h por solicitud."""
        _, today, ts = _tiempos(now)
        req = StudentsService._bloquear(db, request_id)
        try:
            c = account.load(db, req, today)
            if c.total is None:
                raise ServiceError("sin_plan", "La solicitud aún no tiene plan de pagos.", status=409)
            if c.saldo <= 0:
                raise ServiceError("sin_saldo", "La solicitud no tiene saldo pendiente.", status=409)
            if req.last_reminded_at and ts - req.last_reminded_at < _ENTRE_RECORDATORIOS:
                raise ServiceError("ya_recordado", "Ya se le recordó en las últimas 24 horas.", status=409,
                                   siguiente=(req.last_reminded_at + _ENTRE_RECORDATORIOS).isoformat())
            p = c.proxima
            fecha, monto, vencido = fecha_larga(p.due_date), pesos(p.amount - p.cubierto), p.due_date < today
            req.last_reminded_at = ts
            log_event(db, event_type="recordatorio_enviado", request_id=req.id, actor_id=actor_id,
                      payload={"saldo": account.money(c.saldo)})
            db.commit()
        except Exception:
            db.rollback()
            raise
        db.refresh(req)
        body = (f"Saldo {pesos(c.saldo)}. Tienes un pago vencido desde el {fecha} por {monto}." if vencido
                else f"Saldo {pesos(c.saldo)}. Próximo pago: {fecha} por {monto}.")
        notify.notify_student(db, req, kind="recordatorio", link="/prorrogas/estudiante", body=body,
                              context={"saldo": pesos(c.saldo), "proxima_fecha": fecha, "proxima_monto": monto,
                                       "vencido": vencido})
        try:
            db.commit()
        except Exception:
            logger.exception("No se pudo guardar el aviso de la solicitud %s", req.id)
            db.rollback()
        return req

    @staticmethod
    def export_adeudos(db, period_id: int, actor_id: int, *, now: date | datetime | None = None) -> tuple[bytes, str]:
        """Lista de estudiantes con parcialidad vencida (§16.1) para el centro de cómputo."""
        now, today, ts = _tiempos(now)
        period = db.get(AcademicPeriod, period_id)
        if period is None:
            raise ServiceError("periodo_no_encontrado", "No se encontró el semestre.", status=404)
        reqs = (db.query(ProrrogaRequest).filter(ProrrogaRequest.period_id == period_id,
                                                 ProrrogaRequest.status == "aceptada")
                .order_by(ProrrogaRequest.student_name, ProrrogaRequest.id).all())
        cuentas = account.load_many(db, reqs, today)
        filas = [(r, cuentas[r.id]) for r in reqs if cuentas[r.id].situacion == "atrasada"]

        ws = Workbook().active
        ws.title = "Adeudos"
        ws["A1"], ws["A2"] = "Lista de adeudos de prórrogas", f"Semestre: {period_label(period)}"
        ws["A3"] = f"Exportado el {ts:%d/%m/%Y %H:%M}"
        ws["A4"] = f"Criterio: estudiantes con al menos una parcialidad vencida y no cubierta al {today:%d/%m/%Y}."
        ws["A5"] = f"Total de estudiantes: {len(filas)}"
        ws["A1"].font = Font(bold=True, size=14)
        for i, (titulo, ancho) in enumerate(_COLUMNAS, start=1):
            celda = ws.cell(row=7, column=i, value=titulo)
            celda.font = Font(bold=True)
            ws.column_dimensions[celda.column_letter].width = ancho
        for r, c in filas:
            ws.append([r.control_number, r.student_name, r.program_name, r.email, folio(r.id), c.total, c.pagado,
                       c.saldo, sum(1 for p in c.parcialidades if p.estado == "vencida"), c.vencida_desde])
            fila = ws.max_row
            for col in range(1, 6):  # texto: openpyxl haría fórmula de un nombre que empiece con "="
                celda = ws.cell(row=fila, column=col)
                if isinstance(celda.value, str):
                    celda.data_type = "s"
            for col in (6, 7, 8):
                ws.cell(row=fila, column=col).number_format = _DINERO
            ws.cell(row=fila, column=10).number_format = "DD/MM/YYYY"
            # Se conserva la primera exportación; si ya se avisó su desbloqueo, empieza un ciclo nuevo.
            if r.block_exported_at is None or r.unblock_notified_at is not None:
                r.block_exported_at, r.unblock_notified_at = ts, None
        if not filas:
            ws.append(["Sin estudiantes con adeudo"])
        ws.freeze_panes = "A8"
        ws.auto_filter.ref = f"A7:J{max(ws.max_row, 8)}"
        log_event(db, event_type="lista_bloqueo_exportada", request_id=None, actor_id=actor_id,
                  payload={"period_id": period_id, "estudiantes": len(filas), "solicitudes": [r.id for r, _ in filas]})
        try:
            db.commit()
        except Exception:
            db.rollback()
            raise
        buf = BytesIO()
        ws.parent.save(buf)
        return buf.getvalue(), f"adeudos-prorrogas-{ts:%Y%m%d-%H%M}.xlsx"

    @staticmethod
    def mark_unblock(db, request_id: int, actor_id: int, *, now: date | datetime | None = None) -> ProrrogaRequest:
        """Marca que el estudiante, ya liquidado, fue reportado para desbloqueo (§16.2)."""
        _, today, ts = _tiempos(now)
        req = StudentsService._bloquear(db, request_id)
        try:
            if not _avisar_desbloqueo(req, account.load(db, req, today)):
                raise ServiceError("no_aplica", "No hay un desbloqueo pendiente de avisar.", status=409)
            req.unblock_notified_at = ts
            log_event(db, event_type="desbloqueo_avisado", request_id=req.id, actor_id=actor_id, payload={})
            db.commit()
        except Exception:
            db.rollback()
            raise
        return req
