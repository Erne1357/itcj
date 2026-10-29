"""Llenado de plantillas ``.docx`` con marcadores ``{{nombre}}``.

Word parte el texto en *runs* a su antojo: basta con que la secretaria corrija
una letra para que ``{{nombre}}`` quede como ``{{nom`` + ``bre}}`` en dos runs
distintos. Por eso nada aquí busca dentro de un run suelto: cada párrafo se lee
como el texto concatenado de sus runs, los marcadores se ubican sobre ese texto
y el reemplazo se reparte de vuelta entre los runs. El texto nuevo hereda el
formato del run donde empezaba el marcador.

Se recorren todos los ``w:p`` del cuerpo (tablas y cuadros de texto incluidos)
y de los encabezados y pies propios de cada sección.
"""
from __future__ import annotations

import re
from collections.abc import Iterator, Mapping
from copy import deepcopy
from io import BytesIO

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from docx.text.run import Run

PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-z0-9_]+)\s*\}\}")

_W_P = qn("w:p")
_W_R = qn("w:r")
_W_T = qn("w:t")
_XML_SPACE = qn("xml:space")


class InvalidDocx(ValueError):
    """Los bytes no son un documento de Word (.docx) que python-docx pueda abrir."""


class MissingValues(KeyError):
    """La plantilla trae marcadores para los que no se dio valor."""

    def __init__(self, names):
        self.names = frozenset(names)
        super().__init__(", ".join(sorted(self.names)))


# ---------------------------------------------------------------------------
# Recorrido
# ---------------------------------------------------------------------------

def _open(data: bytes) -> DocxDocument:
    try:
        return docx.Document(BytesIO(data))
    except Exception as exc:  # zip roto, no es OOXML, falta word/document.xml…
        raise InvalidDocx("El archivo no es un documento de Word válido.") from exc


def _save(doc: DocxDocument) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def iter_paragraphs(doc: DocxDocument) -> Iterator[Paragraph]:
    """Todos los párrafos: cuerpo, tablas, cuadros de texto, encabezados y pies."""
    for p_el in doc.element.body.iter(_W_P):
        yield Paragraph(p_el, doc._body)
    seen: set[int] = set()
    for section in doc.sections:
        for hf in (section.header, section.footer,
                   section.first_page_header, section.first_page_footer,
                   section.even_page_header, section.even_page_footer):
            # Ligado al anterior = no tiene parte propia; leerlo la crearía.
            if hf.is_linked_to_previous:
                continue
            root = hf._element
            if id(root) in seen:
                continue
            seen.add(id(root))
            for p_el in root.iter(_W_P):
                yield Paragraph(p_el, hf)


def _owner_paragraph(el):
    parent = el.getparent()
    while parent is not None and parent.tag != _W_P:
        parent = parent.getparent()
    return parent


def _runs_of(p_el) -> list:
    """Runs del párrafo, sin los de párrafos anidados (cuadros de texto)."""
    return [r for r in p_el.iter(_W_R) if _owner_paragraph(r) is p_el]


def _run_text(r) -> str:
    return "".join(t.text or "" for t in r.findall(_W_T))


def _paragraph_text(p_el) -> str:
    return "".join(_run_text(r) for r in _runs_of(p_el))


# ---------------------------------------------------------------------------
# Reemplazo de tramos
# ---------------------------------------------------------------------------

def _replace_spans(paragraph: Paragraph, spans: list[tuple[int, int, str, bytes | None]],
                   image_height_cm: float) -> None:
    """Aplica ``(inicio, fin, texto, imagen)`` sobre el texto concatenado del párrafo.

    Los tramos no se traslapan. Se aplican de atrás hacia adelante para que los
    desplazamientos de los anteriores sigan siendo válidos.
    """
    runs = _runs_of(paragraph._p)
    texts = [_run_text(r) for r in runs]
    original = list(texts)
    starts, acc = [], 0
    for t in texts:
        starts.append(acc)
        acc += len(t)

    def locate_start(pos: int) -> tuple[int, int]:
        for i, (s, t) in enumerate(zip(starts, original)):
            if s <= pos < s + len(t):
                return i, pos - s
        raise IndexError(pos)

    def locate_end(pos: int) -> tuple[int, int]:
        for i, (s, t) in enumerate(zip(starts, original)):
            if s < pos <= s + len(t):
                return i, pos - s
        raise IndexError(pos)

    pictures: list[tuple[int, str, bytes]] = []  # (run, texto que va después, png)
    for start, end, repl, image in sorted(spans, key=lambda s: s[0], reverse=True):
        rs, os_ = locate_start(start)
        re_, oe = locate_end(end)
        if rs == re_:
            tail = texts[rs][oe:]
            if image:
                texts[rs] = texts[rs][:os_]
                pictures.append((rs, tail, image))
            else:
                texts[rs] = texts[rs][:os_] + repl + tail
        else:
            texts[rs] = texts[rs][:os_] + ("" if image else repl)
            for k in range(rs + 1, re_):
                texts[k] = ""
            texts[re_] = texts[re_][oe:]
            if image:
                # Lo que sigue al marcador ya vive en el run `re_`: la imagen va
                # al final del run `rs` y queda en su lugar.
                pictures.append((rs, "", image))

    for r, new, old in zip(runs, texts, original):
        if new == old:
            continue
        ts = r.findall(_W_T)
        ts[0].text = new
        ts[0].set(_XML_SPACE, "preserve")
        for extra in ts[1:]:
            r.remove(extra)

    for idx, tail, image in pictures:
        run_el = runs[idx]
        if tail:
            tail_el = deepcopy(run_el)
            ts = tail_el.findall(_W_T)
            ts[0].text = tail
            for extra in ts[1:]:
                tail_el.remove(extra)
            run_el.addnext(tail_el)
        Run(run_el, paragraph).add_picture(BytesIO(image), height=Cm(image_height_cm))


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

def find_placeholders(docx_bytes: bytes) -> set[str]:
    doc = _open(docx_bytes)
    found: set[str] = set()
    for p in iter_paragraphs(doc):
        found.update(m.group(1) for m in PLACEHOLDER_RE.finditer(_paragraph_text(p._p)))
    return found


def fill_docx(docx_bytes: bytes, values: Mapping[str, str],
              images: Mapping[str, bytes] | None = None, *,
              image_height_cm: float = 1.6) -> bytes:
    """Sustituye cada ``{{marcador}}``; los de ``images`` por la imagen.

    Lanza ``MissingValues`` si algún marcador no tiene valor ni imagen. Un
    marcador de imagen con bytes vacíos, o que solo viene en ``values``, se
    trata como texto (típicamente ``""``: firma pendiente).
    """
    images = {k: v for k, v in (images or {}).items() if v}
    doc = _open(docx_bytes)
    paragraphs = list(iter_paragraphs(doc))

    found: set[str] = set()
    for p in paragraphs:
        found.update(m.group(1) for m in PLACEHOLDER_RE.finditer(_paragraph_text(p._p)))
    missing = found - set(values) - set(images)
    if missing:
        raise MissingValues(missing)

    for p in paragraphs:
        spans = []
        for m in PLACEHOLDER_RE.finditer(_paragraph_text(p._p)):
            name = m.group(1)
            if name in images:
                spans.append((m.start(), m.end(), "", images[name]))
            else:
                spans.append((m.start(), m.end(), str(values[name]), None))
        if spans:
            _replace_spans(p, spans, image_height_cm)
    return _save(doc)


def replace_literals(docx_bytes: bytes, mapping: Mapping[str, str]) -> tuple[bytes, dict[str, int]]:
    """Reemplaza textos literales (en orden) y cuenta cuántas veces apareció cada uno."""
    doc = _open(docx_bytes)
    counts: dict[str, int] = {}
    for old, new in mapping.items():
        counts[old] = 0
        if not old:
            continue
        for p in iter_paragraphs(doc):
            text = _paragraph_text(p._p)
            spans, pos = [], text.find(old)
            while pos != -1:
                spans.append((pos, pos + len(old), new, None))
                pos = text.find(old, pos + len(old))
            if spans:
                counts[old] += len(spans)
                _replace_spans(p, spans, 0)
    return _save(doc), counts
