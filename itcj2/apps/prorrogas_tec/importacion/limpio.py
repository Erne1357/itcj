"""Archivo limpio: lo que una persona revisa antes de importar, y su lectura estricta.

``leer_limpio`` es la compuerta: solo lo que valida aquí puede importarse. Una solicitud
(``ref``) con cualquier ``problema`` se rechaza completa, nunca a medias.
"""
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from ..services.errors import ServiceError
from ..services.formatting import parse_monto
from .hojas import Solicitud

ENCABEZADOS = ("ref", "semestre", "fecha_solicitud", "no_control", "nombre", "carrera", "carrera_original",
               "numero", "vencimiento", "monto", "pagado", "problema", "valor_original", "origen")
_REVISAR = ("hoja", "renglon", "columna", "valor", "motivo")
_IDX = {c: i for i, c in enumerate(ENCABEZADOS)}
_FMT_FECHA, _FMT_MONTO = "DD/MM/YYYY", '"$"#,##0.00'
_MAX_PARCIALIDADES = 6


@dataclass
class SolicitudLimpia:
    ref: str
    semestre: str
    fecha_solicitud: date
    no_control: str
    nombre: str
    carrera: str
    parcialidades: list[tuple[int, date, Decimal, bool]]


def _fila(s: Solicitud, p, numero, extra: str = "") -> list:
    problema = "; ".join(x for x in (s.problema, p.problema if p else extra) if x)
    return [s.ref, s.semestre, s.fecha_solicitud, s.no_control, s.nombre, s.carrera, s.carrera_original,
            numero, p.vencimiento if p else None, float(p.monto) if p and p.monto is not None else None,
            None if not p or p.pagado is None else int(p.pagado), problema,
            p.valor_original if p else "", p.origen if p else ""]


def _revisar(s: Solicitud) -> list[tuple]:
    out = []
    if s.problema:
        valor = (s.no_control if s.problema.startswith("número de control")
                 else s.carrera_original if s.problema.startswith("carrera sin equivalencia") else "")
        out.append((s.hoja, int(s.ref.rsplit("#", 1)[1]), "", valor, s.problema))
    for p in s.parcialidades:
        if p.problema:
            hoja, celda = p.origen.rsplit("!", 1)
            col, fila = re.fullmatch(r"([A-Z]+)(\d+)", celda).groups()
            out.append((hoja, int(fila), col, p.valor_original, p.problema))
    return out


def escribir_limpio(solicitudes: list[Solicitud], path) -> dict:
    wb = Workbook()
    ws, rev = wb.active, wb.create_sheet("Revisar")
    ws.title = "Parcialidades"
    ws.append(list(ENCABEZADOS))
    rev.append(list(_REVISAR))
    n_parc = 0
    for s in solicitudes:
        # sort estable: empates conservan el orden de aparición; sin vencimiento al final
        orden = sorted(s.parcialidades, key=lambda p: (p.vencimiento is None, p.vencimiento or date.min))
        if orden:
            for n, p in enumerate(orden, 1):
                ws.append(_fila(s, p, n))
            n_parc += len(orden)
        else:
            ws.append(_fila(s, None, None, extra="" if s.problema else "sin parcialidades"))
        for r in _revisar(s) or ([(s.hoja, int(s.ref.rsplit("#", 1)[1]), "", "", "sin parcialidades")]
                                 if not s.parcialidades else []):
            rev.append(list(r))
    for row in ws.iter_rows(min_row=2):
        row[_IDX["fecha_solicitud"]].number_format = _FMT_FECHA
        row[_IDX["vencimiento"]].number_format = _FMT_FECHA
        row[_IDX["monto"]].number_format = _FMT_MONTO
    for hoja in (ws, rev):
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
    wb.save(path)
    return {"solicitudes": len(solicitudes), "parcialidades": n_parc,
            "con_problema": sum(1 for s in solicitudes if s.tiene_problema or not s.parcialidades)}


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        try:
            return datetime.strptime(v.strip(), "%d/%m/%Y").date()
        except ValueError:
            return None
    return None


def _texto(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "" if v is None else str(v).strip()


def _validar(ref: str, filas: list[tuple]) -> tuple[SolicitudLimpia | None, str]:
    col = lambda r, c: r[_IDX[c]]  # noqa: E731
    for r in filas:
        if _texto(col(r, "problema")):
            return None, f"tiene datos por revisar: {_texto(col(r, 'problema'))}"
    est = {(_texto(col(r, "semestre")), _fecha(col(r, "fecha_solicitud")) or _texto(col(r, "fecha_solicitud")),
            _texto(col(r, "no_control")), _texto(col(r, "nombre")), _texto(col(r, "carrera"))) for r in filas}
    if len(est) > 1:
        return None, "datos del estudiante distintos entre renglones"
    semestre, fecha, control, nombre, carrera = est.pop()
    if not isinstance(fecha, date):
        return None, "sin fecha de solicitud"
    numeros = [col(r, "numero") for r in filas]
    if all(_texto(n) == "" for n in numeros):
        return None, "sin parcialidades"
    try:
        nums = [int(n) if _texto(n) == _texto(int(n)) else None for n in numeros]
    except (TypeError, ValueError):
        nums = [None]
    if nums != list(range(1, len(filas) + 1)) or len(filas) > _MAX_PARCIALIDADES:
        return None, "numeración de parcialidades inválida"
    venc = [_fecha(col(r, "vencimiento")) for r in filas]
    if None in venc:
        return None, "vencimiento inválido"
    try:
        montos = [parse_monto(_texto(col(r, "monto")) if not isinstance(col(r, "monto"), float)
                              else str(col(r, "monto"))) for r in filas]
    except ServiceError:
        return None, "monto inválido"
    pag = [col(r, "pagado") for r in filas]
    if any(isinstance(p, bool) or p not in (0, 1) or isinstance(p, str) for p in pag):
        return None, "pagado inválido"
    control = control.upper()
    if not re.fullmatch(r"[A-Z]?\d{8}", control):
        return None, "número de control inválido"
    return SolicitudLimpia(ref, semestre, fecha, control, nombre, carrera,
                           list(zip(nums, venc, montos, [bool(p) for p in pag]))), ""


def leer_limpio(path) -> tuple[list[SolicitudLimpia], list[tuple[str, str]]]:
    wb = load_workbook(path, data_only=True)
    if "Parcialidades" not in wb.sheetnames:
        raise ValueError("Falta la hoja Parcialidades")
    filas = list(wb["Parcialidades"].iter_rows(values_only=True))
    if tuple(filas[0] if filas else ()) != ENCABEZADOS:
        raise ValueError(f"Encabezados esperados: {', '.join(ENCABEZADOS)}")
    grupos: dict[str, list[tuple]] = {}
    for n, r in enumerate(filas[1:], 2):
        if _texto(r[0]):
            grupos.setdefault(str(r[0]).strip(), []).append(r)
        elif any(_texto(c) for c in r):
            # no sabemos de qué solicitud era: seguir dejaría una solicitud con parcialidades de menos
            raise ValueError(f"El renglón {n} de Parcialidades tiene datos pero no tiene ref; "
                             "corrígelo antes de importar.")
    validas, errores = [], []
    for ref, rs in grupos.items():
        ok, motivo = _validar(ref, rs)
        if ok:
            validas.append(ok)
        else:
            errores.append((ref, motivo))
    return validas, errores
