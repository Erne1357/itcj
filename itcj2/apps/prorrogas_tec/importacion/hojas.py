"""Lectura del Excel histórico: una ``Solicitud`` por alumno con sus ``Parcialidad``.

Nada se descarta: lo que no se puede interpretar sin adivinar queda como ``problema``
en la solicitud o en la parcialidad, con el valor original y la celda de origen.
"""
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .normaliza import (
    MESES, NoInterpretable, _vacio, mapear_carrera, parse_bandera, parse_control, parse_fechas, parse_importes,
)

# nombre exacto de la hoja → (semestre, formato)
HOJAS: dict[str, tuple[str, str]] = {
    "Inscripción Enero - Junio 2026": ("enero – junio 2026", "mensual"),
    "Inscripción Agosto-Diciembre 25": ("agosto – diciembre 2025", "mensual"),
    "Inscripción Enero-Junio 2025": ("enero – junio 2025", "mensual"),
    "Prórrogas agosto-dic 2024": ("agosto – diciembre 2024", "mensual"),
    "Inscripción Agosto-Dic 2024": ("agosto – diciembre 2024", "importe_unico"),
    "Enero Junio 2024": ("enero – junio 2024", "celda"),
}

_RE_ITEM = re.compile(r"(\d{1,2}\s+de\s+[a-záéíóú]+(?:\s+(?:de\s+)?\d{4})?)\s*(\$\s?[\d.,]+)", re.IGNORECASE)


@dataclass
class Parcialidad:
    vencimiento: date | None
    monto: Decimal | None
    pagado: bool | None
    problema: str = ""
    valor_original: str = ""
    origen: str = field(kw_only=True)  # "{hoja}!{celda}"


@dataclass
class Solicitud:
    ref: str  # "{hoja}#{renglón}"
    hoja: str
    semestre: str
    fecha_solicitud: date | None
    no_control: str
    nombre: str
    carrera: str
    carrera_original: str
    parcialidades: list[Parcialidad]
    problema: str = ""  # problema de la solicitud entera

    @property
    def tiene_problema(self) -> bool:
        return bool(self.problema) or any(p.problema for p in self.parcialidades)


def _texto(v) -> str:
    return "" if _vacio(v) else str(v)


def _problema(motivo: str, valor: str, origen: str) -> Parcialidad:
    return Parcialidad(None, None, None, motivo, valor, origen=origen)


def _celda(valor, origen: str, anio: int) -> list[Parcialidad]:
    """Formato ``celda``: ``30 de enero $2,900.00`` (uno o varios); cualquier falla marca la celda entera."""
    texto = str(valor)
    items = _RE_ITEM.findall(texto)
    try:
        if not items or _RE_ITEM.sub("", texto).strip():
            raise NoInterpretable("fecha ilegible")
        res = []
        for f, i in items:
            (fecha,), (monto,) = parse_fechas(f, anio=anio), parse_importes(i)  # el regex garantiza uno y uno
            res.append(Parcialidad(fecha, monto, True, origen=origen))
        return res
    except NoInterpretable as e:
        return [_problema(e.motivo, texto, origen)]


def _grupo(vf, vi, vb, origen: str, anio: int) -> list[Parcialidad]:
    """Terna fecha · importe · bandera; fechas e importes de la misma celda se emparejan en orden."""
    try:
        fechas, montos, pagado = parse_fechas(vf, anio=anio), parse_importes(vi), parse_bandera(vb)
        if not fechas and not montos:
            return []  # una bandera 0/1 sola no dice nada; una ilegible sí cae en el except
        if fechas and len(fechas) == len(montos):
            return [Parcialidad(f, m, pagado, origen=origen) for f, m in zip(fechas, montos)]
        motivo = ("parcialidad sin importe" if not montos else "parcialidad sin fecha" if not fechas
                  else "fechas e importes no cuadran")
    except NoInterpretable as e:
        motivo = e.motivo
    return [_problema(motivo, f"{_texto(vf)} | {_texto(vi)} | {_texto(vb)}", origen)]


def _leer_hoja(ws) -> list[Solicitud]:
    hoja = ws.title
    if hoja not in HOJAS:
        return [Solicitud(f"{hoja}#0", hoja, "", None, "", "", "", "", [], "hoja desconocida")]
    semestre, formato = HOJAS[hoja]
    anio = int(semestre.split()[-1])
    filas = ws.iter_rows(values_only=True)
    enc = ["" if h is None else str(h).strip() for h in next(filas, ())]
    col = {h: i for i, h in enumerate(enc) if h}

    # grupos de columnas (fecha, importe, bandera) ubicados por encabezado
    grupos, ok = [], True
    if formato == "importe_unico":
        ok = "FECHA" in col and "IMPORTE" in col
        grupos = [(col["FECHA"], col["IMPORTE"], None)] if ok else []
    else:
        for i, h in enumerate(enc):
            if h.lower() not in MESES:
                continue
            if formato == "celda":
                grupos.append((i, None, None))
            elif enc[i + 1:i + 2] == ["IMPORTE"]:
                grupos.append((i, i + 1, i + 2 if enc[i + 2:i + 3] == [""] else None))
            else:
                ok = False
    if not ok or not {"No. De Control", "NOMBRE", "CARRERA"} <= col.keys():
        return [Solicitud(f"{hoja}#1", hoja, semestre, None, "", "", "", "", [], "encabezados no reconocidos")]

    c_no = col.get("No.")
    de_grupo = {c for g in grupos for c in g if c is not None}
    de_alumno = {col[k] for k in ("FECHA", "No. De Control", "NOMBRE", "CARRERA") if k in col}
    if formato == "celda":
        de_alumno.discard(col.get("FECHA"))  # el formato celda no tiene fecha de solicitud

    sols: list[Solicitud] = []
    validas: list[Solicitud] = []
    actual = None
    for r, fila in enumerate(filas, start=2):
        if all(_vacio(v) for i, v in enumerate(fila) if i != c_no):
            actual = None  # un hueco corta la continuación: lo de abajo no se pega al alumno de arriba
            continue
        celda = {k: fila[i] for k, i in col.items()}
        nueva = not _vacio(celda.get("No. De Control"))
        # en una continuación, lo que traiga en las columnas del alumno va a revisión
        leidas = de_grupo | de_alumno if nueva else de_grupo
        pagos = []
        for cf, ci, cb in grupos:
            origen = f"{hoja}!{get_column_letter(cf + 1)}{r}"
            vf, vi, vb = (None if c is None else fila[c] for c in (cf, ci, cb))
            if formato == "celda":
                pagos += [] if _vacio(vf) else _celda(vf, origen, anio)
            else:
                pagos += _grupo(vf, vi, vb, origen, anio)
        pagos += [_problema("dato sin interpretar", str(v), f"{hoja}!{get_column_letter(i + 1)}{r}")
                  for i, v in enumerate(fila) if i not in leidas and i != c_no and not _vacio(v)]
        if nueva:
            actual = _nueva(hoja, semestre, formato, anio, r, celda, validas)
            sols.append(actual)
        elif actual is None and pagos:  # un renglón con solo banderas 0/1 no abre nada
            actual = Solicitud(f"{hoja}#{r}", hoja, semestre, None, "", "", "", "", [], "renglón sin estudiante")
            sols.append(actual)
        if actual is not None:
            actual.parcialidades += pagos

    repetidos = Counter(s.no_control for s in validas)
    for s in validas:
        if repetidos[s.no_control] > 1 and not s.problema:
            s.problema = "número de control repetido en la hoja"
    if formato == "celda":
        for s in sols:
            s.fecha_solicitud = min((p.vencimiento for p in s.parcialidades if p.vencimiento), default=None)
    return sols


def _nueva(hoja, semestre, formato, anio, r, celda: dict, validas: list) -> Solicitud:
    """Renglón con número de control: valida control, carrera y fecha de solicitud (gana el primer problema)."""
    problemas = []
    try:
        no_control, valido = parse_control(celda["No. De Control"]), True
    except NoInterpretable as e:
        no_control, valido = str(celda["No. De Control"]), False
        problemas.append(e.motivo)
    carrera_original = _texto(celda.get("CARRERA"))
    try:
        carrera = mapear_carrera(carrera_original)
    except NoInterpretable as e:
        carrera = ""
        problemas.append(e.motivo)
    fecha = None
    if formato != "celda":
        try:
            fechas = parse_fechas(celda.get("FECHA"), anio=anio)
        except NoInterpretable:
            fechas = []
        if len(fechas) == 1:
            fecha = fechas[0]
        else:
            problemas.append("fecha de solicitud ilegible")
    s = Solicitud(f"{hoja}#{r}", hoja, semestre, fecha, no_control, " ".join(_texto(celda.get("NOMBRE")).split()),
                  carrera, carrera_original, [], problemas[0] if problemas else "")
    if valido:
        validas.append(s)
    return s


def leer_libro(path) -> list[Solicitud]:
    """Todas las solicitudes del libro, en el orden de las hojas y los renglones."""
    wb = load_workbook(path, data_only=True)
    return [s for ws in wb.worksheets for s in _leer_hoja(ws)]
